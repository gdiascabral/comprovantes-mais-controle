# -*- coding: utf-8 -*-
"""Conta do ERP que ainda não tem linha no painel: achar e incluir.

Até aqui, conta nova no ERP aparecia só como AVISO no fim do resumo ("conta
nova no ERP fora do painel") e os pagamentos dela ficavam fora da planilha,
somados no log das contas fora do painel. Pôr a conta no painel exigia mexer
em TRÊS arquivos à mão, e combinados: uma linha nova no `MODELO.xlsx` (com as
fórmulas de todas as colunas), a entrada dela no `mapping.yaml` e a faixa de
linhas do `config.yaml`. Errar um deles não dá erro no Excel — dá saldo na
linha de outra conta, ou um aporte que some da soma. Aqui os três mudam
juntos, ou nenhum muda.

A régua de "fora do painel" é a MESMA do aviso (`rules.resolve_balances`): a
lista que a pessoa vê no botão é exatamente a que o resumo acusa. Uma
segunda régua seria a chance de o botão dizer "nada novo" com o aviso na
tela.

**A linha nova entra no FIM, depois da última conta.** Inserir no meio
desceria as linhas de todas as contas de baixo, e o `mapping.yaml` guarda o
NÚMERO da linha de cada uma — trocar a ordem é trabalho do Excel e do mapa
juntos, e não precisa acontecer para a conta nova funcionar. No fim, as
linhas antigas ficam onde estavam e só o que vem depois delas desce (os
TOTAIS DO DIA e o «Como ler»).

**A linha nova é uma cópia da última conta, como o Excel copia**: estilo e
fórmulas de cada coluna, com as referências relativas andando uma linha
(`$B33` vira `$B34`) e as absolutas paradas (`Movimentações!$C$54:$C$73`).
Quem faz isso é o `Translator` do próprio openpyxl, que já vai inteiro no exe
(`--collect-all openpyxl`). O que não é fórmula na última linha (os valores
do dia, o "de quem vem") nasce vazio.

**As faixas que terminam na última conta esticam**, em qualquer aba, e o que
está abaixo dela desce — é o que o Excel faz ao inserir linhas logo abaixo de
uma faixa, mais o esticar que ele NÃO faz sozinho: `=SUM(C8:C33)` passa a
somar a linha nova, o `VLOOKUP(…, Painel!$B$8:$H$33, …)` da aba
«Movimentações» passa a enxergá-la, a formatação condicional (linha âmbar,
saldo vermelho) passa a pintá-la e a área de impressão passa a incluí-la.
Faixa que termina ANTES da última conta é escolha de quem desenhou o modelo,
e fica como está.

**Nada é trocado sem prova** (`_provar`): os três arquivos novos são escritos
ao lado dos de verdade e relidos pelo MESMO código que gera o painel do dia —
o `config.yaml` e o `mapping.yaml` recarregam, a coluna B bate com o mapa
(`check_labels`), toda linha antiga continua casando com as MESMAS contas do
ERP, cada linha nova casa com a sua e só com ela, e uma planilha de prova é
montada pelo `workbook.build`, que confere o modelo célula por célula. Passou,
os três de antes vão para uma cópia datada e os novos tomam o lugar; falhou
alguma troca no meio (o Excel com o modelo aberto é o caso comum), o que já
tinha sido trocado volta da cópia.
"""
from __future__ import annotations

import json
import os
import re
import shutil
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange

from .config import PlanilhaConfig, load_config
from .mapping import AccountMapping
from .models import ErpAccount
from .parsing import extract_account_numbers, normalize_account_number, normalize_name
from .rules import build_row_fills, resolve_balances
from .workbook import WorkbookError, build, check_labels

ARQ_CONFIG = "config.yaml"
ARQ_MAPA = "mapping.yaml"
ARQ_MODELO = "MODELO.xlsx"

#: Onde ficam os três arquivos de antes de cada inclusão, uma pasta datada
#: por vez, ao lado deles. Voltar atrás é copiar os três de volta.
PASTA_COPIAS = "copias do painel"

#: O SUMIF da aba «Movimentações» usa a coluna B como CRITÉRIO, e critério do
#: Excel não é texto cru: `*`, `?` e `~` são curinga, e `<`, `>` e `=` no
#: começo viram comparação. Um nome assim somaria o aporte de outra conta, ou
#: de nenhuma, sem erro na tela. E texto que começa com `=` o openpyxl grava
#: como FÓRMULA.
_CURINGAS = ("*", "?", "~")
_COMECO_PROIBIDO = ("=", "+", "-", "@", "<", ">")

#: O Excel não aceita critério de SUMIF com mais de 255 caracteres.
_MAX_ROTULO = 255


class InclusaoRecusada(Exception):
    """Nada foi gravado — e a mensagem diz por quê."""


@dataclass(frozen=True)
class Inclusao:
    """Uma conta do ERP e o nome que a linha dela terá na coluna B."""

    conta: ErpAccount
    rotulo: str


@dataclass(frozen=True)
class ResultadoInclusao:
    #: (número da linha no painel, nome na coluna B), na ordem em que entraram.
    linhas: list[tuple[int, str]]
    #: A pasta com os três arquivos como estavam antes.
    copia: Path


# --------------------------------------------------------------------------
# Achar
# --------------------------------------------------------------------------
def contas_fora_do_painel(contas: list[ErpAccount],
                          mapping: AccountMapping) -> list[ErpAccount]:
    """As contas do ERP que não casam com linha viva do painel nem estão na
    lista de ignoradas, em ordem de nome."""
    fora = resolve_balances(contas, mapping).contas_desconhecidas
    return sorted(fora, key=lambda c: normalize_name(c.name))


def rotulo_sugerido(conta: ErpAccount) -> str:
    """O nome da linha, para ser corrigido e não digitado: o nome da conta no
    ERP, com os espaços dobrados desfeitos."""
    return re.sub(r"\s+", " ", conta.name or "").strip()


def problemas_da_inclusao(inclusoes: list[Inclusao],
                          mapping: AccountMapping) -> list[str]:
    """Tudo o que impede incluir, de uma vez. Lista vazia = pode incluir.

    Junta TODOS os problemas, e não o primeiro: quem vai corrigir quer a lista
    inteira, e não descobrir um nome por clique.
    """
    if not inclusoes:
        return ["nenhuma conta marcada."]
    problemas: list[str] = []
    # O SUMIF do Excel não distingue maiúscula de minúscula: dois nomes que
    # só diferem na caixa somariam o mesmo aporte nas duas linhas.
    existentes = {r.label.strip().casefold(): r.row for r in mapping.rows}
    vistos: set[str] = set()
    ids: set[str] = set()
    for inc in inclusoes:
        nome = inc.conta.name
        rotulo = (inc.rotulo or "").strip()
        linha = mapping.resolve_account(inc.conta)
        if linha is not None and linha.exists_in_erp:
            problemas.append(f"{nome}: já está no painel (linha {linha.row}).")
        if inc.conta.id in ids:
            problemas.append(f"{nome}: marcada duas vezes.")
        ids.add(inc.conta.id)
        if not rotulo:
            problemas.append(f"{nome}: falta o nome da linha.")
            continue
        if rotulo.startswith(_COMECO_PROIBIDO):
            problemas.append(
                f"{rotulo}: o nome não pode começar com "
                f"{' '.join(_COMECO_PROIBIDO)} — o Excel lê como fórmula ou "
                "comparação.")
        if any(c in rotulo for c in _CURINGAS):
            problemas.append(
                f"{rotulo}: tire {' '.join(_CURINGAS)} do nome — na soma da "
                "aba Movimentações o Excel lê esses sinais como curinga.")
        if len(rotulo) > _MAX_ROTULO:
            problemas.append(
                f"{rotulo[:40]}…: nome com mais de {_MAX_ROTULO} letras — o "
                "Excel não soma por um nome desse tamanho.")
        chave = rotulo.casefold()
        if chave in existentes:
            problemas.append(
                f"{rotulo}: já existe uma linha com esse nome "
                f"(linha {existentes[chave]}).")
        elif chave in vistos:
            problemas.append(f"{rotulo}: nome repetido entre as marcadas.")
        vistos.add(chave)
    return problemas


# --------------------------------------------------------------------------
# O MODELO.xlsx
# --------------------------------------------------------------------------
_REF_CELULA = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d+)$")


def _nova_faixa(r1: int, r2: int, ultima: int, n: int) -> tuple[int, int]:
    """A regra de TODAS as referências, de linha de fórmula a formatação.

    Inteira abaixo da última conta: desce `n`. Terminando na última conta ou
    passando dela: estica `n` (é a faixa das contas, ou a que as contém).
    Terminando antes: fica."""
    if r1 > ultima:
        return r1 + n, r2 + n
    if r2 >= ultima:
        return r1, r2 + n
    return r1, r2


def _ajustar_ref(ref: str, ultima: int, n: int) -> str:
    """`C33`, `$B$8:$H$33`… com as linhas refeitas. O que não for célula ou
    faixa de células (coluna inteira, nome definido) volta como veio."""
    partes = ref.split(":")
    achados = [_REF_CELULA.match(p) for p in partes]
    if len(partes) not in (1, 2) or not all(achados):
        return ref
    linhas = [int(m.group(4)) for m in achados]
    if len(linhas) == 1:
        novas = [linhas[0] + n if linhas[0] > ultima else linhas[0]]
    else:
        novas = list(_nova_faixa(min(linhas), max(linhas), ultima, n))
        if linhas[0] > linhas[1]:
            novas.reverse()
    return ":".join(f"{m.group(1)}{m.group(2)}{m.group(3)}{nova}"
                    for m, nova in zip(achados, novas))


def _ajustar_formula(formula, aba: str, na_propria_aba: bool,
                     ultima: int, n: int):
    """Refaz as referências ao PAINEL dentro de uma fórmula.

    Na própria aba do painel valem as referências sem nome de aba; nas outras,
    só as que dizem `Painel!`. `Movimentações!$B$9:$B$73`, escrito no painel,
    fala de outra aba e fica como está. Quem separa operando de texto é o
    `Tokenizer` do openpyxl — expressão regular solta confundiria
    `"…C33…"` dentro de aspas com uma célula."""
    if not (isinstance(formula, str) and formula.startswith("=")):
        return formula
    tok = Tokenizer(formula)
    mudou = False
    for t in tok.items:
        if t.type != Token.OPERAND or t.subtype != Token.RANGE:
            continue
        valor = t.value
        if "!" in valor:
            folha, ref = valor.rsplit("!", 1)
            if folha.strip("'").replace("''", "'") != aba:
                continue
            novo = f"{folha}!{_ajustar_ref(ref, ultima, n)}"
        elif na_propria_aba:
            novo = _ajustar_ref(valor, ultima, n)
        else:
            continue
        if novo != valor:
            t.value = novo
            mudou = True
    return tok.render() if mudou else formula


def _ajustar_multi(sqref, ultima: int, n: int) -> str:
    faixas = []
    for faixa in MultiCellRange(str(sqref)).ranges:
        r1, r2 = _nova_faixa(faixa.min_row, faixa.max_row, ultima, n)
        faixas.append(CellRange(min_col=faixa.min_col, min_row=r1,
                                max_col=faixa.max_col, max_row=r2).coord)
    return " ".join(faixas)


def _recusar_o_que_nao_sei_mover(ws, ultima: int) -> None:
    """Tabela do Excel e link abaixo das contas não descem por este código.

    O modelo de hoje não tem nenhum dos dois; se um dia tiver, é melhor
    recusar a inclusão com o motivo do que gravar um modelo torto."""
    if ws.tables:
        raise WorkbookError(
            f"a aba {ws.title!r} tem tabela do Excel ({', '.join(ws.tables)}) "
            "— a inclusão automática não sabe movê-la. Inclua a linha à mão.")
    for linha in ws.iter_rows(min_row=ultima + 1):
        for cel in linha:
            if cel.hyperlink is not None:
                raise WorkbookError(
                    f"a célula {ws.title}!{cel.coordinate} tem link — a "
                    "inclusão automática não sabe movê-lo. Inclua a linha à "
                    "mão.")


def acrescentar_linhas(wb, planilha: PlanilhaConfig, rotulos: list[str]) -> None:
    """Acrescenta uma linha por rótulo logo depois da última conta. Muda `wb`.

    Três passos, nesta ordem: (1) as referências de TODAS as abas passam a
    falar do painel já crescido; (2) o que está abaixo da última conta desce;
    (3) as linhas novas nascem como cópia da última conta."""
    n = len(rotulos)
    if n == 0:
        return
    if planilha.aba not in wb.sheetnames:
        raise WorkbookError(
            f"o modelo não tem a aba {planilha.aba!r}; encontrei {wb.sheetnames}")
    ws = wb[planilha.aba]
    ultima = planilha.ultima_linha
    _recusar_o_que_nao_sei_mover(ws, ultima)

    # (1) As referências — fórmulas, formatação condicional, validação,
    # nomes definidos e a área de impressão.
    for folha in wb.worksheets:
        propria = folha.title == planilha.aba
        for linha in folha.iter_rows():
            for cel in linha:
                if isinstance(cel.value, str) and cel.value.startswith("="):
                    cel.value = _ajustar_formula(cel.value, planilha.aba,
                                                 propria, ultima, n)
        novas = ConditionalFormattingList()
        for cf in folha.conditional_formatting:
            alvo = _ajustar_multi(cf.sqref, ultima, n) if propria else str(cf.sqref)
            for regra in cf.rules:
                regra.formula = [
                    _ajustar_formula(f"={f}", planilha.aba, propria,
                                     ultima, n)[1:]
                    for f in (regra.formula or [])]
                novas.add(alvo, regra)
        folha.conditional_formatting = novas
        if propria:
            for dv in folha.data_validations.dataValidation:
                dv.sqref = MultiCellRange(_ajustar_multi(dv.sqref, ultima, n))
        for nome in list(folha.defined_names.values()):
            if isinstance(nome.attr_text, str):
                nome.attr_text = _ajustar_formula(
                    f"={nome.attr_text}", planilha.aba, propria, ultima, n)[1:]
    for nome in list(wb.defined_names.values()):
        if isinstance(nome.attr_text, str):
            nome.attr_text = _ajustar_formula(
                f"={nome.attr_text}", planilha.aba, False, ultima, n)[1:]
    if ws.print_area:
        area = str(ws.print_area).rsplit("!", 1)[-1].replace("$", "")
        ws.print_area = _ajustar_ref(area, ultima, n)

    # (2) O que está abaixo da última conta desce — de baixo para cima, para
    # nenhuma linha ser escrita antes de ter sido lida. Mesclagens saem antes
    # e voltam depois, no lugar novo.
    mescladas = [m.coord for m in ws.merged_cells.ranges if m.min_row > ultima]
    for coord in mescladas:
        ws.unmerge_cells(coord)
    colunas = ws.max_column
    for r in range(ws.max_row, ultima, -1):
        for c in range(1, colunas + 1):
            origem, destino = ws.cell(r, c), ws.cell(r + n, c)
            destino.value = origem.value
            destino._style = copy(origem._style)
            destino.comment = origem.comment
            origem.comment = None
        dim_o, dim_d = ws.row_dimensions[r], ws.row_dimensions[r + n]
        dim_d.height, dim_d.hidden = dim_o.height, dim_o.hidden
    for coord in mescladas:
        faixa = CellRange(coord)
        faixa.shift(row_shift=n)
        ws.merge_cells(faixa.coord)

    # (3) As linhas novas: cópia da última conta.
    for k, rotulo in enumerate(rotulos, start=1):
        r = ultima + k
        for c in range(1, colunas + 1):
            modelo, cel = ws.cell(ultima, c), ws.cell(r, c)
            cel._style = copy(modelo._style)
            cel.comment = None
            valor = modelo.value
            if isinstance(valor, str) and valor.startswith("="):
                cel.value = Translator(valor, origin=modelo.coordinate
                                       ).translate_formula(cel.coordinate)
            else:
                cel.value = None
        ws.cell(r, 2).value = rotulo
        dim_u, dim_r = ws.row_dimensions[ultima], ws.row_dimensions[r]
        dim_r.height, dim_r.hidden = dim_u.height, dim_u.hidden


# --------------------------------------------------------------------------
# O mapping.yaml e o config.yaml — editados como TEXTO
# --------------------------------------------------------------------------
# Os dois são comentados à mão, e o comentário é metade do valor deles
# (explica por que a linha 17 não tem conta, por que uma conta é ignorada).
# `yaml.safe_dump` apagaria tudo; por isso a edição é textual, como a do
# `erp/discover.aplicar_uuids`, e o resultado é relido pelo carregador de
# verdade antes de ser aceito.

def _yaml_str(texto: str) -> str:
    """Escalar YAML entre aspas duplas. O escape do JSON é um subconjunto do
    YAML de aspas duplas, e aspas protegem de `:`, `#` e número puro (um
    número de conta sem aspas viraria inteiro)."""
    return json.dumps(str(texto), ensure_ascii=False)


def entradas_do_mapa(inclusoes: list[Inclusao], primeira_linha: int,
                     mapping: AccountMapping) -> list[dict]:
    """O que cada linha nova leva no `mapping.yaml`.

    O `uuid` sempre — é a chave que sobrevive a renomear a conta no ERP. O
    nome do ERP sempre — é por ele que os PAGAMENTOS casam, porque a lista de
    pagamentos traz o nome da conta e não o id. O número da conta só quando é
    inequívoco (o que o ERP informa ou o único escrito no nome) e nenhuma
    outra linha já o usa: número repetido no mapa derruba a carga inteira."""
    usados = {r.account_number for r in mapping.rows if r.account_number}
    entradas = []
    for k, inc in enumerate(inclusoes):
        numero = normalize_account_number(inc.conta.account_number)
        if not numero:
            achados = extract_account_numbers(inc.conta.name)
            numero = achados[0] if len(achados) == 1 else None
        if numero in usados:
            numero = None
        if numero:
            usados.add(numero)
        entradas.append({"row": primeira_linha + k,
                         "label": inc.rotulo.strip(),
                         "uuid": inc.conta.id or None,
                         "account_number": numero,
                         "erp_name": inc.conta.name})
    return entradas


def acrescentar_no_mapping(texto: str, entradas: list[dict]) -> str:
    """Põe as entradas no fim da lista `model_rows`, com o recuo que ela já
    usa, e deixa todo o resto do arquivo — comentários inclusive — igual."""
    linhas = texto.splitlines()
    inicio = next((i for i, lin in enumerate(linhas)
                   if re.match(r"^model_rows:\s*(#.*)?$", lin)), None)
    if inicio is None:
        raise InclusaoRecusada(
            "o mapping.yaml não tem a lista 'model_rows' — não sei onde "
            "acrescentar a conta.")
    fim = next((i for i in range(inicio + 1, len(linhas))
                if re.match(r"^[A-Za-z_][\w-]*\s*:", linhas[i])), len(linhas))
    # Comentário e linha em branco logo antes da chave seguinte são DELA.
    while fim > inicio + 1 and (not linhas[fim - 1].strip()
                                or linhas[fim - 1].lstrip().startswith("#")):
        fim -= 1
    recuo = next((m.group(1) for lin in linhas[inicio + 1:fim]
                  if (m := re.match(r"^(\s*)-\s*row\s*:", lin))), "")
    bloco = []
    for e in entradas:
        bloco.append(f"{recuo}- row: {int(e['row'])}")
        bloco.append(f"{recuo}  label: {_yaml_str(e['label'])}")
        if e.get("uuid"):
            bloco.append(f"{recuo}  uuid: {_yaml_str(e['uuid'])}")
        if e.get("account_number"):
            bloco.append(f"{recuo}  account_number: {_yaml_str(e['account_number'])}")
        bloco.append(f"{recuo}  erp_name: {_yaml_str(e['erp_name'])}")
    return "\n".join(linhas[:fim] + bloco + linhas[fim:]) + "\n"


def atualizar_config(texto: str, n: int) -> str:
    """Desce `n` linhas a última conta, a linha dos totais e a célula do total
    dos pagamentos, sem tocar em mais nada do arquivo."""
    def somar(padrao, formatar, rotulo):
        nonlocal texto
        achados = re.findall(padrao, texto, flags=re.M)
        if len(achados) != 1:
            raise InclusaoRecusada(
                f"o config.yaml deveria ter UM '{rotulo}' e tem {len(achados)} "
                "— não sei qual mudar.")
        texto = re.sub(padrao, formatar, texto, count=1, flags=re.M)

    somar(r"^(\s+ultima_linha:\s*)(\d+)",
          lambda m: f"{m.group(1)}{int(m.group(2)) + n}", "ultima_linha")
    somar(r"^(\s+linha_totais:\s*)(\d+)",
          lambda m: f"{m.group(1)}{int(m.group(2)) + n}", "linha_totais")
    somar(r"^(\s+celula_total_pagamentos:\s*[\"']?)([A-Za-z]+)(\d+)",
          lambda m: f"{m.group(1)}{m.group(2)}{int(m.group(3)) + n}",
          "celula_total_pagamentos")
    return texto


# --------------------------------------------------------------------------
# Incluir: montar ao lado, provar, copiar, trocar
# --------------------------------------------------------------------------
def _casamento(contas: list[ErpAccount], mapping: AccountMapping) -> dict:
    """linha do painel -> ids das contas do ERP que casam com ela."""
    casadas: dict[int, list[str]] = {}
    for conta in contas:
        if mapping.is_ignored(conta.name):
            continue
        linha = mapping.resolve_account(conta)
        if linha is not None and linha.exists_in_erp:
            casadas.setdefault(linha.row, []).append(conta.id)
    return casadas


def _provar(modelo: Path, mapa: Path, config: Path, antigo: AccountMapping,
            antiga_ultima: int, contas_erp: list[ErpAccount],
            inclusoes: list[Inclusao], prova: Path) -> None:
    """Relê os três arquivos novos com o código do dia. Levanta se algo não
    fecha — e aí nada foi trocado."""
    cfg = load_config(config)
    novo = AccountMapping.load(mapa)
    pl = cfg.planilha
    n = len(inclusoes)
    problemas = []
    if pl.ultima_linha != antiga_ultima + n:
        problemas.append(f"o config.yaml ficou com a última conta na linha "
                         f"{pl.ultima_linha}, e não na {antiga_ultima + n}.")
    wb = openpyxl.load_workbook(modelo)
    try:
        check_labels(wb[pl.aba], novo)
    finally:
        wb.close()
    antes, depois = _casamento(contas_erp, antigo), _casamento(contas_erp, novo)
    for linha, ids in sorted(antes.items()):
        if depois.get(linha) != ids:
            problemas.append(f"a linha {linha} deixaria de casar com a mesma "
                             "conta do ERP.")
    for k, inc in enumerate(inclusoes, start=1):
        ids = depois.get(antiga_ultima + k)
        if ids != [inc.conta.id]:
            outras = len(ids or []) - 1
            problemas.append(
                f"{inc.conta.name}: a linha nova "
                + ("não casaria com a conta." if not ids else
                   f"casaria também com mais {outras} conta(s) do ERP — o "
                   "nome é parecido demais com outra."))
    if problemas:
        raise InclusaoRecusada("\n".join(problemas))
    # A planilha de prova é montada pelo mesmo `build` do dia, que confere o
    # modelo inteiro depois de gravar — e é apagada em seguida.
    try:
        build(modelo, prova, date.today(), build_row_fills(novo, {}, {}),
              novo, pl)
    finally:
        try:
            prova.unlink()
        except OSError:
            pass


def incluir_no_painel(pasta, inclusoes: list[Inclusao],
                      contas_erp: list[ErpAccount], *,
                      agora: datetime | None = None) -> ResultadoInclusao:
    """Inclui as contas no painel: `MODELO.xlsx`, `mapping.yaml` e
    `config.yaml`, os três de uma vez ou nenhum.

    `contas_erp` é a lista inteira de contas do ERP lida agora — é contra ela
    que se prova que as linhas antigas continuam casando com as mesmas
    contas. Levanta `InclusaoRecusada` com o motivo; nesse caso os arquivos
    ficaram como estavam."""
    pasta = Path(pasta)
    cam_cfg, cam_mapa = pasta / ARQ_CONFIG, pasta / ARQ_MAPA
    cfg = load_config(cam_cfg)
    cam_modelo = (cfg.caminho("modelo") if "modelo" in cfg.caminhos
                  else pasta / ARQ_MODELO)
    mapping = AccountMapping.load(cam_mapa)
    inclusoes = [Inclusao(i.conta, (i.rotulo or "").strip()) for i in inclusoes]
    problemas = problemas_da_inclusao(inclusoes, mapping)
    if problemas:
        raise InclusaoRecusada("\n".join(problemas))

    pl = cfg.planilha
    n = len(inclusoes)
    primeira = pl.ultima_linha + 1
    novos = {cam_modelo: cam_modelo.with_name(f"{cam_modelo.stem}.novo.xlsx"),
             cam_mapa: cam_mapa.with_name(f"{cam_mapa.stem}.novo.yaml"),
             cam_cfg: cam_cfg.with_name(f"{cam_cfg.stem}.novo.yaml")}
    prova = cam_modelo.with_name(f"{cam_modelo.stem}.prova.xlsx")
    try:
        try:
            wb = openpyxl.load_workbook(cam_modelo)
            try:
                acrescentar_linhas(wb, pl, [i.rotulo for i in inclusoes])
                wb.save(novos[cam_modelo])
            finally:
                wb.close()
            novos[cam_mapa].write_text(
                acrescentar_no_mapping(
                    cam_mapa.read_text(encoding="utf-8"),
                    entradas_do_mapa(inclusoes, primeira, mapping)),
                encoding="utf-8")
            novos[cam_cfg].write_text(
                atualizar_config(cam_cfg.read_text(encoding="utf-8"), n),
                encoding="utf-8")
            _provar(novos[cam_modelo], novos[cam_mapa], novos[cam_cfg], mapping,
                    pl.ultima_linha, contas_erp, inclusoes, prova)
        except InclusaoRecusada:
            raise
        except Exception as e:                              # noqa: BLE001
            # Montar e provar só escrevem nos arquivos `.novo`: qualquer
            # falha aqui deixa os de verdade como estavam, e é isso que a
            # mensagem precisa dizer.
            raise InclusaoRecusada(
                f"não consegui montar o painel novo, e nada foi trocado: {e}"
            ) from e

        copia = pasta / PASTA_COPIAS / f"{(agora or datetime.now()):%Y-%m-%d %H%M%S}"
        copia.mkdir(parents=True, exist_ok=False)
        for final in novos:
            shutil.copy2(final, copia / final.name)
        trocados: list[Path] = []
        try:
            for final, novo in novos.items():
                os.replace(novo, final)
                trocados.append(final)
        except OSError as e:
            for final in trocados:
                shutil.copy2(copia / final.name, final)
            raise InclusaoRecusada(
                f"não consegui gravar {final.name} ({e.strerror or e}). Se ele "
                "está aberto no Excel, feche e tente de novo. Nada mudou: o "
                "que já tinha sido trocado voltou da cópia.") from e
    finally:
        for novo in novos.values():
            try:
                novo.unlink()
            except OSError:
                pass

    return ResultadoInclusao(
        linhas=[(primeira + k, i.rotulo) for k, i in enumerate(inclusoes)],
        copia=copia)
