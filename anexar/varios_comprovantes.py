# -*- coding: utf-8 -*-
"""Anexar N comprovantes num pagamento só — a fase das somas da aba "Anexar (novo)".

Sem tela e sem navegador: a `api` (a `mc_api.MCApi` da sessão) entra por
parâmetro. Quem chama roda na thread do navegador (`AnexarFrame.submeter`).

Por que NÃO usa a função de anexar da aba Anexar (revisão de dinheiro):
- a tela, plano B dela, devolve "ja_tinha" assim que o sub-pagamento tem
  QUALQUER anexo (`mc_client.MCClient`): o 2º comprovante sairia "ok" sem
  subir;
- `mc_api.anexar_por_api` dá "ja_anexado" e prova a subida também pelo
  TAMANHO (`mc_api._mesmo_arquivo`): dois comprovantes de mesmo tamanho fariam
  o 2º "já estar lá".
Por isso: só API; tamanhos distintos; pagamento SEM anexo antes; para no
primeiro erro (POST nunca se repete); prova final pelo NOME de cada arquivo.

Falha no meio nunca vira "anexado": o resultado diz quais arquivos ENTRARAM e
quais FALTARAM. Numa rodada nova, o pagamento já tem anexo e nada sobe de novo
(`erro:ja_tem_anexo`, com a mesma conta de quais dos N já estão lá) — o resto
se completa à mão, conferindo o lançamento.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import util

from . import config, matcher, matcher_somado, mc_api

_OK = (mc_api.ANEXADO, "dry_run")


@dataclass
class ResultadoVarios:
    estado: str
    por_arquivo: list = field(default_factory=list)   # (nome, resposta da API)
    entraram: list = field(default_factory=list)      # nomes que estão no pagamento
    faltaram: list = field(default_factory=list)      # nomes que NÃO estão

    @property
    def ok(self) -> bool:
        return self.estado in _OK


def _listagem(api, paid_id):
    """(lista, "") ou (None, "erro:listagem:…"). Falha ao listar NUNCA vira
    "sem anexo" (-1 não é 0)."""
    try:
        lista = api._anexos_brutos(paid_id)
    except Exception as e:                                   # noqa: BLE001
        return None, f"erro:listagem:{e.__class__.__name__}"
    if not isinstance(lista, list):
        motivo = lista.get("__erro") if isinstance(lista, dict) else type(lista).__name__
        return None, f"erro:listagem:{motivo}"
    return lista, ""


def _provado(lista, arquivo: Path) -> bool:
    """O arquivo aparece na listagem PELO NOME (o tamanho não prova aqui)."""
    nome = arquivo.name.lower()
    return any(isinstance(i, dict) and nome in mc_api._nomes_do_item(i) for i in lista)


def _separar(lista, arquivos):
    entraram = [a.name for a in arquivos if _provado(lista, a)]
    return entraram, [a.name for a in arquivos if a.name not in entraram]


def anexar_varios(api, paid_id, arquivos, *, dry_run=False, log=None) -> ResultadoVarios:
    log = log or (lambda _m: None)
    arquivos = [Path(a) for a in arquivos]
    if not paid_id:
        return ResultadoVarios("erro:sem_paid_id")
    if len(arquivos) < 2:
        return ResultadoVarios("erro:menos_de_dois_arquivos")
    tamanhos = []
    for a in arquivos:
        try:
            tamanhos.append(a.stat().st_size)
        except OSError as e:
            return ResultadoVarios(f"erro:arquivo:{a.name}:{e.__class__.__name__}")
    if len(set(tamanhos)) < len(tamanhos):
        return ResultadoVarios("erro:tamanhos_iguais")
    antes, erro = _listagem(api, paid_id)
    if erro:
        return ResultadoVarios(erro)
    if antes:
        entraram, faltaram = _separar(antes, arquivos)
        return ResultadoVarios("erro:ja_tem_anexo", entraram=entraram, faltaram=faltaram)
    por_arquivo, n = [], len(arquivos)
    for k, a in enumerate(arquivos, 1):
        r = api.anexar_por_api(paid_id, a, log=log, dry_run=dry_run)
        por_arquivo.append((a.name, r))
        log(f"   {k}/{n} {a.name} -> {r}")
        if r == "dry_run":
            continue
        if r != mc_api.ANEXADO:
            entraram = [nome for nome, resp in por_arquivo if resp == mc_api.ANEXADO]
            faltaram = [b.name for b in arquivos if b.name not in entraram]
            talvez = (f" (o {a.name} pode ter ficado lá sem arquivo)"
                      if r.startswith("erro:upload:") else "")
            return ResultadoVarios(
                f"erro:parcial:{k - 1}/{n} subiram, o {k}º deu {r} — entraram: "
                f"{', '.join(entraram) or 'nenhum'}; faltaram: {', '.join(faltaram)}"
                f"{talvez} — confira o lançamento", por_arquivo, entraram, faltaram)
    if dry_run:
        return ResultadoVarios("dry_run", por_arquivo)
    depois, erro = _listagem(api, paid_id)
    if erro:                         # sem a releitura, nenhum está provado pelo nome
        return ResultadoVarios(f"erro:nao_confirmado:{erro}", por_arquivo, [],
                               [a.name for a in arquivos])
    entraram, faltaram = _separar(depois, arquivos)
    if not faltaram:
        return ResultadoVarios(mc_api.ANEXADO, por_arquivo, entraram, [])
    return ResultadoVarios(f"erro:nao_confirmado:{len(entraram)}/{n}", por_arquivo,
                           entraram, faltaram)


@dataclass
class Somas:
    anexados: list = field(default_factory=list)
    duvidas: list = field(default_factory=list)
    nao_verificados: int = 0
    com_anexo: int = 0

    @property
    def ok(self) -> int:
        return sum(1 for pe in self.anexados if pe.get("resultado") in _OK)


def fase_das_somas(api, pagos, pasta, *, simular, carregar_pdfs=matcher.carregar_pdfs,
                   preencher_origem=None, log=print, cancelar=lambda: False) -> Somas:
    """Depois da fase 1 (o `_t_auto` de sempre): quem ficou SEM PAR por falta de
    PDF do valor é conferido DE NOVO no ERP (só `SEM_ANEXO` segue; "não
    verificado" fica de fora, contado) e casado com N PDFs livres pelo
    `matcher_somado`. `pagos` são TODOS os do período.

    Dúvida de soma é quem tem `motivo_soma` (o `matcher_somado` grava e mantém o
    status "SEM PAR"): vai para `somas.duvidas` com o motivo — inclusive a que
    já vinha marcada de antes, que não é somada de novo."""
    somas = Somas()
    ja_em_duvida = [p for p in pagos if p.get("status") == "SEM PAR" and p.get("motivo_soma")]
    somas.duvidas.extend(ja_em_duvida)
    sem_par = [p for p in pagos if p.get("status") == "SEM PAR" and not p.get("motivo_soma")
               and p.get("motivo_sem_par") == matcher_somado.MOTIVO_NENHUM_PDF]
    if not sem_par or cancelar():
        return somas
    att = api.verificar_anexos([p["paidId"] for p in sem_par], log, cancelar=cancelar)
    if cancelar():
        return somas
    estados = {p["paidId"]: mc_api.estado_anexo(att, p["paidId"]) for p in sem_par}
    livres = [p for p in sem_par if estados[p["paidId"]] == mc_api.SEM_ANEXO]
    somas.nao_verificados = sum(1 for e in estados.values() if e == mc_api.NAO_VERIFICADO)
    somas.com_anexo = sum(1 for e in estados.values() if e == mc_api.COM_ANEXO)
    if not livres:
        return somas
    pasta = Path(pasta)
    usados = {p["pdf"] for p in pagos if p.get("status") == "CERTEZA" and p.get("pdf")}
    pdfs = [pd for pd in carregar_pdfs(pasta, log) if pd["fn"] not in usados]
    if preencher_origem is not None:
        try:
            preencher_origem(livres, pdfs)
        except Exception:                                   # noqa: BLE001
            log("[aviso] não deu para saber a conta de origem — as somas seguem sem essa regra.")
    somados, duvidas = matcher_somado.casar_somados(livres, pdfs, todos=pagos)
    somas.duvidas.extend(duvidas)
    for pe in somados:
        if cancelar():
            break
        r = anexar_varios(api, pe["paidId"], [pasta / fn for fn in pe["pdfs"]],
                          dry_run=simular, log=log)
        pe["resultado"], pe["pdf"] = r.estado, " + ".join(pe["pdfs"])
        pe["entraram"], pe["faltaram"] = r.entraram, r.faltaram
        somas.anexados.append(pe)
    return somas


def juntar_fim(fim1: tuple, somas: Somas) -> tuple:
    """O ("fim", …) da fase 1 com as somas: anexados entram no total; os que a
    fase 2 tocou saem do "sem par" (dúvida de soma vira dúvida)."""
    ok, tot, duv, sp, saida = fim1
    n, d = len(somas.anexados), len(somas.duvidas)
    return (ok + somas.ok, tot + n, duv + d, max(sp - n - d, 0), saida)


def _detalhe(p) -> str:
    """Num anexo que não deu certo, quais dos N entraram e quais faltaram."""
    if p.get("resultado") in _OK or not (p.get("entraram") or p.get("faltaram")):
        return ""
    return (f" · entraram: {', '.join(p.get('entraram') or []) or 'nenhum'}"
            f" · faltaram: {', '.join(p.get('faltaram') or [])}")


def resumo_das_somas(somas: Somas) -> list[str]:
    linhas = [f"Somas: {len(somas.anexados)} pagamento(s) com N comprovantes "
              f"({somas.ok} ok) · {len(somas.duvidas)} em dúvida"
              + (f" · {somas.nao_verificados} não verificado(s), ficam para a mão"
                 if somas.nao_verificados else "")
              + (f" · {somas.com_anexo} já com anexo, não mexidos"
                 if somas.com_anexo else "")]
    linhas += [f"  {util.fmt_val(p['valor'])} = {p['pdf']} -> {p['resultado']}{_detalhe(p)}"
               for p in somas.anexados]
    linhas += [f"  {util.fmt_val(p['valor'])} DÚVIDA: {p.get('motivo_soma', '')}"
               for p in somas.duvidas]
    return linhas


_CABECALHO = ["Valor", "Data", "Favorecido", "Centro de custo", "Conta", "Descrição",
              "Nº doc", "PDFs", "Motivo", "Resultado", "Entraram", "Faltaram", "Link"]


def acrescentar_ao_relatorio(caminho, somas: Somas) -> None:
    """Duas abas no Excel que a fase 1 gravou: ANEXADOS POR SOMA e SOMA EM DUVIDA."""
    if not caminho or not (somas.anexados or somas.duvidas):
        return
    from openpyxl import load_workbook              # só quem gera paga o openpyxl
    wb = load_workbook(caminho)

    def comuns(p):
        return [util.fmt_val(p["valor"]), p.get("dataFull", ""), p.get("favorecido", ""),
                "; ".join(p.get("works") or []), p.get("conta", ""), p.get("desc", ""),
                p.get("doc", "")]

    for nome, linhas in (
            ("ANEXADOS POR SOMA", [comuns(p) + [" + ".join(p.get("pdfs") or []),
                                                p.get("motivo", ""), p.get("resultado", ""),
                                                ", ".join(p.get("entraram") or []),
                                                ", ".join(p.get("faltaram") or []),
                                                config.MC_URL_LANCAMENTO + str(p["launchId"])]
                                   for p in somas.anexados]),
            ("SOMA EM DUVIDA", [comuns(p) + ["", p.get("motivo_soma", ""), "", "", "",
                                             config.MC_URL_LANCAMENTO + str(p["launchId"])]
                                for p in somas.duvidas])):
        ws = wb.create_sheet(nome)
        ws.append(_CABECALHO)
        for linha in linhas:
            ws.append(linha)
        ws.freeze_panes = "A2"
    wb.save(caminho)
