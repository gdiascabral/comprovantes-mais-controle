# -*- coding: utf-8 -*-
"""Passo 4 da tela do dia: o MODELO.xlsx do dia com os números da TELA.

Escreve o que o painel antigo escrevia (data, C saldo do dia, D a pagar, J/K)
e mais o que antes era digitado à mão: o APORTE DEFINIDO (Painel H), o VALOR
DEFINIDO de cada aportador (Movimentações H) e os movimentos «À MÃO» (seção 3
da Movimentações, linhas 54 a 73: B origem, C destino, H valor). Só valor
literal — nenhuma fórmula digitada —, e o `workbook.build` confere o resto
célula por célula.

A linha da «Movimentações» de cada regra sai da FÓRMULA da coluna B
(`Regras!$C<n>`): no modelo real há linhas vazias e um título entre as seções,
e "linha da regra + 4" erraria a partir da seção 2.

Participante de uma conta com aporte e sem valor leva 0: célula vazia faria
`I = IF(H<>"",H,G)` cair no VALOR SUGERIDO e somar dinheiro que ninguém definiu.

O Excel cópia tem de bater com a tela: os movimentos à mão entram na seção 3
porque as fórmulas E/F/G/M do Painel já somam essa seção (a tela espelha isso
em `aportes.montar_aportes(a_mao=...)`). Como lá, valor zero ou negativo é
ignorado. Mais de 20 movimentos não cabem: erro, nunca truncar em silêncio.

Leitura interrompida ou ausente nunca vira Excel: seria gravar "a pagar 0".
"""
from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path

import openpyxl

from conciliacao.rules import build_row_fills
from conciliacao.workbook import (A_MAO_LINHAS, ABA_MOVIMENTACOES, BuildResult,
                                  WorkbookError, build)

COLUNA_DEFINIDO = "H"
CABECALHO_MOVIMENTACOES = "H7"
PRIMEIRA_LINHA_A_MAO, ULTIMA_LINHA_A_MAO = A_MAO_LINHAS
MAX_A_MAO = ULTIMA_LINHA_A_MAO - PRIMEIRA_LINHA_A_MAO + 1
ZERO = Decimal("0.00")
_REGRA_NA_FORMULA = re.compile(r"Regras!\$?C\$?(\d+)", re.IGNORECASE)


def linhas_da_movimentacao(ws) -> dict[int, int]:
    mapa = {}
    for (celula,) in ws.iter_rows(min_col=2, max_col=2):
        m = _REGRA_NA_FORMULA.search(str(celula.value or ""))
        if m:
            mapa[int(m.group(1))] = celula.row
    return mapa


def _conferir_cabecalho(ws, coord: str, oque: str) -> None:
    if "DEFINIDO" not in str(ws[coord].value or "").upper():
        raise WorkbookError(
            f"{ws.title}!{coord} devia ser o {oque} e é {ws[coord].value!r}: o modelo "
            "mudou — não escrevo valor numa coluna que não sei qual é.")


def _celulas_a_mao(a_mao, rotulos_do_painel, distribuicoes=()) -> dict[str, object]:
    """B/C/H54..73 da seção «À MÃO». Ignora valor <= 0 (como a tela).

    As distribuições de lucros (item A) vêm DEPOIS dos movimentos, com a conta
    de origem em B e o sócio em C: o E do Painel (SUMIF da origem) as soma, e o
    M (SUMIF do destino) não acha linha do painel com o nome do sócio."""
    movimentos = []
    for origem, destino, valor in (a_mao or ()):
        if not isinstance(valor, Decimal):
            raise WorkbookError(f"movimento à mão com valor que não é Decimal: {valor!r}")
        if valor <= ZERO:
            continue
        if not isinstance(destino, str) or not destino.strip():
            raise WorkbookError("movimento à mão sem conta de destino")
        if origem is not None and (not isinstance(origem, str) or not origem.strip()):
            raise WorkbookError("movimento à mão com origem vazia (use None para 'de fora')")
        if destino not in rotulos_do_painel:
            raise WorkbookError(f"movimento à mão para {destino}, que não é conta do "
                                "painel: a tela o ignora e o Excel divergiria dela.")
        movimentos.append((origem, destino, valor))
    for origem, socio, valor in (distribuicoes or ()):
        if not isinstance(valor, Decimal):
            raise WorkbookError(f"distribuição com valor que não é Decimal: {valor!r}")
        if valor <= ZERO:
            continue
        if not isinstance(socio, str) or not socio.strip():
            raise WorkbookError("distribuição sem o sócio que recebe")
        if origem not in rotulos_do_painel:
            raise WorkbookError(f"distribuição de {origem}, que não é conta do painel: a "
                                "tela a ignora e o Excel divergiria dela.")
        movimentos.append((origem, socio, valor))
    if len(movimentos) > MAX_A_MAO:
        raise WorkbookError(
            f"{len(movimentos)} movimentos à mão (com as distribuições): a seção "
            "«À MÃO» do modelo só tem "
            f"{MAX_A_MAO} linhas ({PRIMEIRA_LINHA_A_MAO} a {ULTIMA_LINHA_A_MAO}).")
    celulas: dict[str, object] = {}
    for i, (origem, destino, valor) in enumerate(movimentos):
        linha = PRIMEIRA_LINHA_A_MAO + i
        if origem is not None:
            celulas[f"B{linha}"] = origem
        celulas[f"C{linha}"] = destino
        celulas[f"{COLUNA_DEFINIDO}{linha}"] = float(valor)
    return celulas


def celulas_de_aporte(modelo, planilha, mapping, regras, definidos, a_mao=(),
                      distribuicoes=()) -> dict:
    validos = {par: valor for par, valor in definidos.items() if valor}
    for (destino, aportador), valor in validos.items():
        if valor < 0:
            raise WorkbookError(f"aporte negativo de {aportador} para {destino}: {valor}")
    mao = _celulas_a_mao(a_mao, {r.label for r in mapping.rows}, distribuicoes)
    if not validos and not mao:
        return {}
    wb = openpyxl.load_workbook(Path(modelo))
    try:
        if planilha.aba not in wb.sheetnames or ABA_MOVIMENTACOES not in wb.sheetnames:
            raise WorkbookError(f"o modelo precisa das abas {planilha.aba!r} e "
                                f"{ABA_MOVIMENTACOES!r}; encontrei {wb.sheetnames}")
        _conferir_cabecalho(wb[ABA_MOVIMENTACOES], CABECALHO_MOVIMENTACOES,
                            "VALOR DEFINIDO")
        if validos:
            _conferir_cabecalho(wb[planilha.aba],
                                f"{COLUNA_DEFINIDO}{planilha.primeira_linha - 1}",
                                "APORTE DEFINIDO")
        linhas_mov = linhas_da_movimentacao(wb[ABA_MOVIMENTACOES])
    finally:
        wb.close()

    regras_do_par = {(r.destino, r.aportador) for r in regras}
    linha_do_rotulo = {r.label: r.row for r in mapping.rows}
    totais: dict[str, Decimal] = {}
    vistos: set = set()
    for r in regras:
        par = (r.destino, r.aportador)
        if par in vistos and any(d == r.destino for d, _ in validos):
            raise WorkbookError(f"{r.aportador} → {r.destino} aparece mais de uma vez na "
                                "aba «Regras»: o valor seria contado em dobro no Excel.")
        vistos.add(par)
    for (destino, aportador), valor in validos.items():
        if (destino, aportador) not in regras_do_par:
            raise WorkbookError(f"{aportador} → {destino} não está na aba «Regras» do "
                                "modelo: não sei em que linha da «Movimentações» escrever.")
        totais[destino] = totais.get(destino, ZERO) + valor

    painel, mov = {}, dict(mao)
    for destino, total in totais.items():
        row = linha_do_rotulo.get(destino)
        if row is None or row not in planilha.linhas:
            raise WorkbookError(f"{destino!r} não é uma linha do painel.")
        painel[f"{COLUNA_DEFINIDO}{row}"] = float(total)
        for regra in regras:
            if regra.destino != destino:
                continue
            linha_mov = linhas_mov.get(regra.linha)
            if linha_mov is None:
                raise WorkbookError(f"a regra da linha {regra.linha} da aba «Regras» não "
                                    f"tem linha na «{ABA_MOVIMENTACOES}».")
            mov[f"{COLUNA_DEFINIDO}{linha_mov}"] = float(
                validos.get((destino, regra.aportador), ZERO))
    extras = {}
    if painel:
        extras[planilha.aba] = painel
    if mov:
        extras[ABA_MOVIMENTACOES] = mov
    return extras


def gerar(*, modelo, destino, data, mapping, planilha, saldos, a_pagar, regras,
          definidos, a_mao=(), distribuicoes=(),
          leitura_interrompida: bool = False) -> BuildResult:
    if leitura_interrompida:
        raise WorkbookError("a leitura dos pagamentos foi interrompida: não gero o Excel "
                            "do dia com 'a pagar' incompleto.")
    if a_pagar is None:
        raise WorkbookError("a leitura dos pagamentos está ausente: não gero o Excel do "
                            "dia sem ela (sairia 'a pagar 0').")
    extras = celulas_de_aporte(modelo, planilha, mapping, regras, definidos, a_mao,
                               distribuicoes)
    fills = build_row_fills(mapping, saldos, a_pagar)
    return build(modelo, destino, data, fills, mapping, planilha, extras=extras)
