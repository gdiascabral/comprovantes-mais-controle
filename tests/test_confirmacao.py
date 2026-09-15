# -*- coding: utf-8 -*-
"""A janela "Confirmar o que entra" do passo 2, sem janela nenhuma.

Até 14/09/2026 ela abria ANTES da leitura dos anexos: mostrava "BOLETO sem
código de barras" para o título cujo boleto estava dentro do PDF da nota, e
não mostrava nada do que depois ia para a aba NÃO ENTRARAM. O dono quase
deixou pagamentos de fora por causa disso. Agora ela mostra a leitura real —
o que a planilha vai ter, o que a remessa vai fazer com cada linha e o que
não entrou, com o motivo —, e a regra de tudo isso mora em
`pagamentos_dia/confirmacao.py`, puro, testado aqui.

Nenhum dado real: nomes, chaves, OCs e valores são inventados. O repo é
público.
"""
import datetime as _dt
from pathlib import Path

import pytest

from extratos_sicoob import sicoob_contas
from pagamentos_dia import confirmacao
from pagamentos_dia import regras_pagamento as regras
from pagamentos_dia import relatorio
from pagamentos_dia import remessa_dia
from pagamentos_dia.remessa_dia import Candidato
from relatorios import contas_mc

#: Linha digitável sintética cujos dígitos verificadores FECHAM — a mesma dos
#: testes da remessa. Valor embutido: R$ 1.150,00.
LINHA_BANCARIA = "34191.57007 00024.924375 24177.010006 9 15340000115000"
HOJE = _dt.date(2026, 9, 14)
CONTA = "EMPRESA EXEMPLO - SICOOB"


def _lanc(ident, conta="CONTA A", favorecido="Fornecedor Modelo Ltda",
          valor=100.0, **extra):
    """Um lançamento a pagar, na forma que a API do ERP devolve."""
    item = {"id": ident, "tradePayableId": f"T-{ident}", "paidTo": favorecido,
            "remainingValue": valor, "plannedDate": "2026-09-14",
            "tradePayableAccount": {"name": conta},
            "tradePayablePaymentMethod": "Pix",
            "paidToBankAccount": "PIX EMAIL fornecedor@exemplo.com",
            "documentNumber": "1234",
            "costCentreDetails": [{"workName": "OBRA MODELO"}],
            "paid": False}
    item.update(extra)
    return item


def _entradas(lancamentos, anexos=None, textos=None, regras_fornecedor=None):
    return confirmacao.Entradas(
        selecionados=list(lancamentos), anexos=anexos or {}, overviews={},
        textos=textos or {}, urls_ocr=set(), cadastro_reembolso={},
        regras_fornecedor=regras_fornecedor or {}, participantes={},
        periodo=(HOJE, HOJE))


# ==========================================================================
# O que não entrou precisa carregar o que a janela mostra
# ==========================================================================
def test_o_nao_apto_carrega_id_e_o_que_a_janela_mostra():
    """Sem o `id` a janela não sabe de que lançamento é a linha vermelha, e
    sem OC, centro de custo e observação o dono não acha o título no ERP."""
    item = _lanc("L9", valor=250.0, paidToBankAccount="",
                 documentNumber="", description="Material OC 4321")
    res = relatorio.montar_registros([item], {}, {}, {})
    assert res.contas == {}
    o, = res.omitidos
    assert o["motivo"] == regras.MOTIVO_SEM_PAGAR
    assert (o["id"], o["oc"], o["centro_custo"], o["dados"]) \
        == ("L9", "4321", "OBRA MODELO", "")
    assert "Pix sem chave" in o["obs"]
    assert not o.get("fora_do_recorte")


def test_a_conta_fora_do_recorte_se_anuncia():
    """A janela não mostra a conta ignorada — mas quem decide isso é uma
    chave, e não o começo de um texto que alguém pode reescrever."""
    item = _lanc("L7", conta="EMPRESA - APENAS LANCAMENTO")
    o, = relatorio.montar_registros([item], {}, {}, {}).omitidos
    assert o["fora_do_recorte"] is True
    assert o["id"] == "L7"


def test_a_aba_nao_entraram_continua_com_as_mesmas_colunas(tmp_path):
    """As chaves novas são da janela; o Excel lê só as de sempre."""
    item = _lanc("L9", paidToBankAccount="", documentNumber="")
    res = relatorio.montar_registros([item], {}, {}, {})
    arquivo = relatorio.gerar_excel(res, tmp_path / "p.xlsx", log=lambda *_: None)
    from openpyxl import load_workbook
    ws = load_workbook(arquivo)[relatorio.ABA_OMITIDOS]
    assert [c.value for c in ws[3]] == relatorio.HEADERS_OMITIDOS
    assert len([c for c in ws[4] if c.value not in (None, "")]) \
        == len(relatorio.HEADERS_OMITIDOS)


# ==========================================================================
# A remontagem: o que a pessoa desmarcou sai, sem baixar nada de novo
# ==========================================================================
def test_o_desmarcado_sai_com_o_motivo_e_o_resto_fica():
    e = _entradas([_lanc("L1"), _lanc("L2", favorecido="Outro Fornecedor Ltda")])
    todos = confirmacao.remontar(e)
    assert sorted(r["id"] for r in todos.contas["CONTA A"]) == ["L1", "L2"]
    assert todos.omitidos == []

    sem_l2 = confirmacao.remontar(e, {"L2"})
    assert [r["id"] for r in sem_l2.contas["CONTA A"]] == ["L1"]
    assert [(o["id"], o["motivo"]) for o in sem_l2.omitidos] \
        == [("L2", regras.MOTIVO_NAO_CONFIRMADO)]


def test_a_planilha_sai_do_periodo_da_busca_e_avisa_se_a_tela_mudou():
    """Os lançamentos em memória são os do período BUSCADO; nomear a planilha
    (e o `_periodo_do_resultado`) pelas datas da tela no clique do passo 2
    punha o dia 15 no nome de uma planilha do dia 14."""
    dia14, dia15 = (_dt.date(2026, 9, 14),) * 2, (_dt.date(2026, 9, 15),) * 2
    assert confirmacao.periodo_da_planilha(dia14, dia14) == (dia14, "")
    periodo, aviso = confirmacao.periodo_da_planilha(dia14, dia15)
    assert periodo == dia14
    assert aviso == ("as datas na tela (15/09/2026 a 15/09/2026) não são as da "
                     "busca (14/09/2026 a 14/09/2026): a planilha sai do período "
                     "buscado — para outro período, busque de novo")
    periodo, aviso = confirmacao.periodo_da_planilha(dia14, None)
    assert periodo == dia14 and "14/09/2026 a 14/09/2026" in aviso
    assert confirmacao.periodo_da_planilha(None, dia14) == (None, "")


def test_o_aviso_de_anexos_nao_lidos_diz_quantos():
    assert confirmacao.aviso_de_anexos_nao_lidos(0) == ""
    assert confirmacao.aviso_de_anexos_nao_lidos(2) == (
        "2 anexo(s) não foram lidos: a forma de pagar dessas linhas pode "
        "estar errada")


def test_a_remontagem_usa_os_anexos_ja_lidos():
    """Confirmar não baixa nada: o boleto lido na apuração continua lido.

    Sem os textos guardados, remontar a planilha depois da janela devolveria
    exatamente o "BOLETO sem código de barras" que a janela nova veio
    desmentir."""
    url = "https://exemplo.invalid/boleto-oc-1234.pdf"
    anexo = {"filename": "boleto oc 1234", "tagName": "Boleto",
             "extension": ".pdf", "downloadUrl": url}
    item = _lanc("L1", valor=1150.0, tradePayablePaymentMethod="Boleto",
                 paidToBankAccount="")
    lido = _entradas([item], anexos={"T-L1": [anexo]},
                     textos={url: f"Linha digitável: {LINHA_BANCARIA}"})
    reg, = confirmacao.remontar(lido).contas["CONTA A"]
    assert reg["dados"] == LINHA_BANCARIA

    nao_lido = _entradas([item], anexos={"T-L1": [anexo]})
    reg, = confirmacao.remontar(nao_lido).contas["CONTA A"]
    assert reg["dados"] == "", "sem o texto guardado não há de onde tirar a linha"


# ==========================================================================
# A análise da remessa, na apuração do passo 2
# ==========================================================================
def _registro(ident, **troca):
    """Uma linha da planilha como o `montar_registros` a deixa."""
    base = {"tipo": "Boleto", "dados": LINHA_BANCARIA, "valor": 1150.00,
            "descricao": "OC 1234 - material", "favorecido": "Fornecedor Modelo Ltda",
            "status": "APTO", "conferencia": "", "obs": "", "id": ident,
            "parcial": False, "oc": "1234", "centro_custo": "OBRA MODELO"}
    base.update(troca)
    return base


class _RegistroFalso:
    """O registro de remessas, com as leituras que o `preparar` faz e as
    duas escritas que a análise NUNCA pode fazer — anotadas, para o teste
    cobrar que ficaram vazias."""

    def __init__(self, enviados=None, mudo=False):
        self.enviados, self.mudo, self.escritas = enviados or {}, mudo, []

    def maior_ordem_do_dia(self, _quando):
        if self.mudo:
            raise OSError("sem rede")
        return 0

    def envio_de(self, _codigo):
        return None

    def envio_da_referencia(self, referencia):
        nsa = self.enviados.get(referencia)
        if not nsa:
            return None
        remessa = type("R", (), {"nsa": nsa,
                                 "gerado_em": _dt.date(2026, 9, 10)})()
        return remessa, None

    def alocar_nsa(self, convenio):
        self.escritas.append(("alocar_nsa", convenio))
        return 1

    def registrar(self, *a, **k):
        self.escritas.append(("registrar", a, k))


def _mapas():
    mapa = contas_mc.Mapa(raiz=Path("."), destinos=[
        contas_mc.Destino(erp=CONTA, empresa="EXEMPLO", pasta="SICOOB",
                          banco="SICOOB"),
        contas_mc.Destino(erp="EMPRESA EXEMPLO - INTER", empresa="EXEMPLO",
                          pasta="INTER", banco="INTER"),
    ])
    empresas = [sicoob_contas.Empresa(
        nome="EXEMPLO",
        contas=[sicoob_contas.Conta(numero="12.345-6", pasta="SICOOB",
                                    banco="756", agencia="4321-0",
                                    convenio="123456")],
        cnpj="11.222.333/0001-81", razao_social="EMPRESA EXEMPLO LTDA")]
    return mapa, empresas


def test_a_analise_so_le_e_diz_o_que_a_remessa_faria():
    registro = _RegistroFalso(enviados={"L2": 7})
    contas = {CONTA: [_registro("L1"), _registro("L2"),
                      _registro("L3", parcial=True)],
              "EMPRESA EXEMPLO - INTER": [_registro("L4")]}
    analise = confirmacao.analisar_remessa(
        contas, {}, carregar_mapas=_mapas, abrir_historico=lambda: registro,
        quando=HOJE)
    assert analise.avisos == []
    assert analise.contas_conferidas is True
    assert analise.sem_remessa == {
        "EMPRESA EXEMPLO - INTER": remessa_dia.MOTIVO_FORA_SICOOB}
    assert analise.candidato(CONTA, "L1").pode
    assert analise.candidato(CONTA, "L2").ja_enviado \
        == "já saiu na remessa nº 000007 de 10/09/2026"
    assert analise.candidato(CONTA, "L3").impedimento == remessa_dia.MOTIVO_PARCIAL
    assert analise.candidato(CONTA, "NAO-EXISTE") is None
    assert registro.escritas == [], "conferir não reserva NSA nem registra nada"


def test_sem_registro_de_remessas_a_planilha_segue_com_aviso():
    def sem_sessao():
        raise RuntimeError("ninguém entrou neste computador ainda")

    analise = confirmacao.analisar_remessa(
        {CONTA: [_registro("L1")]}, {}, carregar_mapas=_mapas,
        abrir_historico=sem_sessao, quando=HOJE)
    aviso, = analise.avisos
    assert aviso.startswith(confirmacao.AVISO_SEM_REGISTRO)
    assert "ninguém entrou" in aviso
    assert analise.candidato(CONTA, "L1").pode, "o resto da análise continua"


def test_registro_que_nao_responde_tambem_vira_aviso():
    """O `preparar` levanta `RegistroMudo` quando a nuvem cala — na remessa é
    recusa, porque o "seu número" sairia repetido. Aqui não sai arquivo
    nenhum, e parar a planilha por isso seria esconder o dia inteiro."""
    analise = confirmacao.analisar_remessa(
        {CONTA: [_registro("L1")]}, {}, carregar_mapas=_mapas,
        abrir_historico=lambda: _RegistroFalso(mudo=True), quando=HOJE)
    aviso, = analise.avisos
    assert aviso.startswith(confirmacao.AVISO_SEM_REGISTRO)
    assert analise.candidato(CONTA, "L1").pode


def test_sem_cadastro_de_contas_a_analise_nao_inventa_conta_sem_remessa():
    def torto():
        raise ValueError("linha 3 do contas_mc.json sem pasta")

    analise = confirmacao.analisar_remessa(
        {CONTA: [_registro("L1")]}, {}, carregar_mapas=torto,
        abrir_historico=_RegistroFalso, quando=HOJE)
    aviso, = analise.avisos
    assert aviso.startswith(confirmacao.AVISO_SEM_CADASTRO)
    assert "linha 3" in aviso
    assert analise.contas_conferidas is False
    assert analise.sem_remessa == {}


# ==========================================================================
# "Já saiu em remessa?" em LOTE, só na análise da janela
# ==========================================================================
class _RegistroDeLote(_RegistroFalso):
    """O registro com a pergunta em lote, que anota quem perguntou o quê.

    `envio_de`/`envio_da_referencia` anotam cada pergunta um-a-um: é assim que
    o teste prova que a análise NÃO as fez para as chaves pré-carregadas — e
    que as fez para as outras."""

    def __init__(self, enviados=None, lote_cai=False, lote_esquece=(),
                 mudo=False, cai_em=()):
        super().__init__(enviados=enviados, mudo=mudo)
        self.lote_cai, self.lote_esquece = lote_cai, set(lote_esquece)
        self.cai_em = set(cai_em)
        self.lotes, self.um_a_um = [], []

    def envio_de(self, codigo):
        self.um_a_um.append(("envio_de", codigo))
        return super().envio_de(codigo)

    def envio_da_referencia(self, referencia):
        self.um_a_um.append(("envio_da_referencia", referencia))
        if referencia in self.cai_em:
            raise OSError("a rede caiu no meio da pergunta")
        return super().envio_da_referencia(referencia)

    def envios_em_lote(self, identificadores, referencias):
        self.lotes.append((sorted(identificadores), sorted(referencias)))
        if self.lote_cai:
            raise OSError("o lote caiu no meio")
        por_barras = {c: None for c in identificadores
                      if c not in self.lote_esquece}
        por_ref = {r: _RegistroFalso.envio_da_referencia(self, r)
                   for r in referencias if r not in self.lote_esquece}
        return por_barras, por_ref


def test_a_analise_pergunta_em_lote_e_nao_uma_por_uma():
    registro = _RegistroDeLote(enviados={"L2": 7})
    contas = {CONTA: [_registro("L1"), _registro("L2"),
                      _registro("L3", parcial=True)]}
    analise = confirmacao.analisar_remessa(
        contas, {}, carregar_mapas=_mapas, abrir_historico=lambda: registro,
        quando=HOJE)
    assert analise.candidato(CONTA, "L2").ja_enviado \
        == "já saiu na remessa nº 000007 de 10/09/2026"
    assert analise.candidato(CONTA, "L1").ja_enviado == ""
    assert len(registro.lotes) == 1
    assert registro.um_a_um == [], \
        "tudo o que o preparar pergunta já estava no lote"
    assert registro.escritas == []


def test_o_que_ficou_fora_do_lote_e_perguntado_ao_historico_de_verdade():
    """Nunca "não saiu" sem ter perguntado: chave que o lote não respondeu vai
    ao caminho de sempre, uma por uma."""
    registro = _RegistroDeLote(enviados={"L2": 7}, lote_esquece={"L2"})
    analise = confirmacao.analisar_remessa(
        {CONTA: [_registro("L1"), _registro("L2")]}, {}, carregar_mapas=_mapas,
        abrir_historico=lambda: registro, quando=HOJE)
    assert analise.candidato(CONTA, "L2").ja_enviado \
        == "já saiu na remessa nº 000007 de 10/09/2026"
    assert ("envio_da_referencia", "L2") in registro.um_a_um
    assert ("envio_da_referencia", "L1") not in registro.um_a_um


def test_o_involucro_so_responde_o_que_pre_carregou():
    registro = _RegistroDeLote(enviados={"L2": 7, "L9": 3})
    envolto = confirmacao.HistoricoPreCarregado(registro, ["1" * 44], ["L2"])
    assert envolto.envio_da_referencia("L2")[0].nsa == 7
    assert envolto.envio_de("1" * 44) is None, "perguntado e não saiu"
    assert registro.um_a_um == []
    assert envolto.envio_da_referencia("L9")[0].nsa == 3
    assert envolto.envio_de("2" * 44) is None
    assert registro.um_a_um == [("envio_da_referencia", "L9"),
                                ("envio_de", "2" * 44)], \
        "fora do pré-carregamento, delegou"
    assert envolto.maior_ordem_do_dia(HOJE) == 0, \
        "o resto do histórico passa direto"


def test_lote_que_cai_volta_ao_um_a_um():
    registro = _RegistroDeLote(enviados={"L2": 7}, lote_cai=True)
    analise = confirmacao.analisar_remessa(
        {CONTA: [_registro("L1"), _registro("L2")]}, {}, carregar_mapas=_mapas,
        abrir_historico=lambda: registro, quando=HOJE)
    assert analise.avisos == [], "o lote é atalho; o caminho de sempre respondeu"
    assert analise.candidato(CONTA, "L2").ja_enviado \
        == "já saiu na remessa nº 000007 de 10/09/2026"
    assert ("envio_da_referencia", "L2") in registro.um_a_um


def test_historico_sem_lote_pergunta_um_a_um():
    """O espelho local e os dublês antigos não têm `envios_em_lote`."""
    registro = _RegistroFalso(enviados={"L2": 7})
    envolto = confirmacao.HistoricoPreCarregado(registro, [], ["L2"])
    assert envolto.envio_da_referencia("L2")[0].nsa == 7


def test_as_chaves_pre_carregadas_sao_as_que_o_preparar_pergunta():
    """Deriva igual ao `preparar`: o código de barras do boleto (44 dígitos,
    pela conversão da linha digitável) e o id de toda linha."""
    contas = {CONTA: [_registro("L1"),
                      _registro("L2", tipo="Pix", dados="11.222.333/0001-81"),
                      _registro("", dados="")]}
    identificadores, referencias = confirmacao.chaves_do_preparar(contas)
    from pagamentos_dia import ocr_boleto
    assert identificadores == [ocr_boleto.codigo_de_barras(LINHA_BANCARIA)]
    assert len(identificadores[0]) == 44
    assert referencias == ["L1", "L2"]


def test_cem_linhas_custam_poucas_consultas_ao_banco(monkeypatch):
    """O registro de verdade, sobre um `rest.ler` que só conta: cem linhas
    (metade boleto, metade Pix com CNPJ na chave) passam a custar a ordem do
    dia e um punhado de consultas `in.(…)` — e nenhuma `eq.` uma por uma."""
    from nuvem import registro as registro_nuvem

    filtros = []

    def ler(tabela, _token, *, colunas="*", filtro=""):
        filtros.append(filtro)
        return []

    monkeypatch.setattr(registro_nuvem.rest, "ler", ler)
    linhas = ([_registro(f"B{i:03d}") for i in range(50)]
              + [_registro(f"P{i:03d}", tipo="Pix", dados="11.222.333/0001-81")
                 for i in range(50)])
    analise = confirmacao.analisar_remessa(
        {CONTA: linhas}, {}, carregar_mapas=_mapas,
        abrir_historico=lambda: registro_nuvem.Registro("tok"), quando=HOJE)
    assert analise.avisos == []
    assert all(c.pode for c in analise.preparado[CONTA])
    assert not [f for f in filtros if "=eq." in f], filtros
    assert len(filtros) <= 5, filtros


def test_falha_da_ordem_do_dia_nao_joga_fora_o_que_o_lote_respondeu():
    """A ordem do dia é do "seu número", que a análise descarta. Cair nela
    rodava o `preparar` de novo SEM histórico, e o "já saiu na remessa nº 7"
    que o lote tinha trazido sumia — a linha voltava verde."""
    registro = _RegistroDeLote(enviados={"L2": 7}, mudo=True)
    analise = confirmacao.analisar_remessa(
        {CONTA: [_registro("L1"), _registro("L2")]}, {}, carregar_mapas=_mapas,
        abrir_historico=lambda: registro, quando=HOJE)
    aviso, = analise.avisos
    assert aviso.startswith(confirmacao.AVISO_SEM_REGISTRO)
    assert analise.candidato(CONTA, "L2").ja_enviado \
        == "já saiu na remessa nº 000007 de 10/09/2026"
    assert analise.envio_conferido("L1") and analise.envio_conferido("L2"), \
        "o lote respondeu as duas: não há o que duvidar"


def test_pergunta_fora_do_lote_que_cai_vira_nao_sei_so_dela():
    registro = _RegistroDeLote(lote_esquece={"L2"}, cai_em={"L2"})
    analise = confirmacao.analisar_remessa(
        {CONTA: [_registro("L1"), _registro("L2")]}, {}, carregar_mapas=_mapas,
        abrir_historico=lambda: registro, quando=HOJE)
    assert analise.avisos and analise.avisos[0].startswith(
        confirmacao.AVISO_SEM_REGISTRO)
    assert analise.envio_conferido("L1") is True
    assert analise.envio_conferido("L2") is False
    assert analise.candidato(CONTA, "L2").ja_enviado == ""


def test_sem_registro_a_linha_diz_que_nao_conferiu_e_fica_ambar():
    def sem_nuvem():
        raise RuntimeError("sem internet")

    contas = {CONTA: [_registro("L1")]}
    analise = confirmacao.analisar_remessa(
        contas, {}, carregar_mapas=_mapas, abrir_historico=sem_nuvem,
        quando=HOJE)
    assert analise.envio_conferido("L1") is False
    grupo, = confirmacao.grupos_da_confirmacao(
        relatorio.Resultado(contas, []), analise, [])
    linha, = grupo.entram
    assert confirmacao.NAO_CONFERI_ENVIO in linha.situacao
    assert linha.estado == "atencao"


# ==========================================================================
# SITUAÇÃO: o veredito da planilha e o da remessa, lado a lado
# ==========================================================================
def _cand(ident="L1", **mudancas):
    campos = dict(id=ident, conta_erp="CONTA A", tipo="Boleto", valor=100.0,
                  favorecido="Fornecedor Modelo Ltda", descricao="",
                  status="APTO")
    campos.update(mudancas)
    return Candidato(**campos)


def test_situacao_junta_a_planilha_e_a_remessa():
    assert confirmacao.situacao_da_linha({"status": "APTO"}, _cand()) \
        == ("APTO · vai na remessa", "ok")
    assert confirmacao.situacao_da_linha(
        {"status": "ATENÇÃO — sem anexo"}, _cand(status="ATENÇÃO — sem anexo")) \
        == ("ATENÇÃO — sem anexo · vai na remessa", "atencao")


def test_o_que_a_remessa_nao_leva_diz_o_motivo_da_propria_remessa():
    """O texto é o MOTIVO_* do `remessa_dia`, e não uma segunda redação: dois
    textos para o mesmo motivo divergem, e aí a janela e a conferência da
    remessa passam a mandar conferir coisas diferentes."""
    texto, estado = confirmacao.situacao_da_linha(
        {"status": "APTO"}, _cand(impedimento=remessa_dia.MOTIVO_PARCIAL))
    assert texto == f"APTO · remessa não leva: {remessa_dia.MOTIVO_PARCIAL}"
    assert estado == "atencao"


def test_a_conta_sem_remessa_pesa_em_toda_linha_dela():
    texto, estado = confirmacao.situacao_da_linha(
        {"status": "APTO"}, _cand(),
        sem_remessa=remessa_dia.MOTIVO_SEM_CONVENIO)
    assert texto == f"APTO · conta sem remessa: {remessa_dia.MOTIVO_SEM_CONVENIO}"
    assert estado == "atencao"


def test_conta_de_outro_banco_nao_e_pendencia():
    """Conta do Inter não faz remessa CNAB e nunca vai fazer: pintar toda linha
    dela de âmbar todo dia é ensinar a pular o âmbar. Ela diz, neutra, que a
    conta não faz remessa e se paga pelo HTML, e a cor vem da planilha."""
    texto, estado = confirmacao.situacao_da_linha(
        {"status": "APTO"}, _cand(), sem_remessa=remessa_dia.MOTIVO_FORA_SICOOB)
    assert texto == (f"APTO · {remessa_dia.MOTIVO_FORA_SICOOB} — "
                     f"{confirmacao.PAGUE_PELO_HTML}")
    assert estado == "ok"
    # O ATENÇÃO da planilha continua âmbar.
    assert confirmacao.situacao_da_linha(
        {"status": "ATENÇÃO — sem anexo"}, _cand(),
        sem_remessa=remessa_dia.MOTIVO_FORA_SICOOB)[1] == "atencao"


@pytest.mark.parametrize("motivo", [
    remessa_dia.MOTIVO_MAO,
    remessa_dia.MOTIVO_PARCIAL,
    remessa_dia.MOTIVO_LINHA,
    remessa_dia.MOTIVO_VALOR_DIVERGE,
    remessa_dia.MOTIVO_REEMBOLSO,
    remessa_dia.MOTIVO_CHAVE_AMBIGUA,
])
def test_em_conta_de_outro_banco_o_impedimento_de_quem_paga_a_mao_e_ambar(motivo):
    """A conta do Inter se paga À MÃO, pelo HTML — e estes impedimentos são
    justamente o que quem paga à mão precisa ver: pagar a outra pessoa, não
    pagar boleto pela metade, linha que não fecha, valor que diverge, reembolso
    sem saber quem recebe, e a chave de onze dígitos que tanto pode ser CPF
    quanto celular — quem a digita no app do banco precisa saber disso. A
    situação diz o motivo, e não o recado genérico."""
    assert motivo in confirmacao.MOTIVOS_DE_QUEM_PAGA_A_MAO
    texto, estado = confirmacao.situacao_da_linha(
        {"status": "APTO"}, _cand(impedimento=motivo),
        sem_remessa=remessa_dia.MOTIVO_FORA_SICOOB)
    assert texto == f"APTO · {motivo}"
    assert confirmacao.PAGUE_PELO_HTML not in texto
    assert estado == "atencao"


def test_so_estes_motivos_sao_de_quem_paga_a_mao():
    """A lista referencia as constantes do `remessa_dia` — não copia texto."""
    assert confirmacao.MOTIVOS_DE_QUEM_PAGA_A_MAO == (
        remessa_dia.MOTIVO_MAO, remessa_dia.MOTIVO_PARCIAL,
        remessa_dia.MOTIVO_LINHA, remessa_dia.MOTIVO_VALOR_DIVERGE,
        remessa_dia.MOTIVO_REEMBOLSO, remessa_dia.MOTIVO_CHAVE_AMBIGUA)


@pytest.mark.parametrize("motivo", [
    remessa_dia.MOTIVO_SEM_DOCUMENTO,
    remessa_dia.MOTIVO_COPIA_COLA,
    remessa_dia.MOTIVO_SANESC,
])
def test_em_conta_de_outro_banco_o_impedimento_tecnico_da_remessa_fica_neutro(
        motivo):
    """Sem CPF/CNPJ para o segmento B, copia-e-cola, SANESC: são coisas do
    ARQUIVO do banco, e esta conta não gera arquivo nenhum."""
    texto, estado = confirmacao.situacao_da_linha(
        {"status": "APTO"}, _cand(impedimento=motivo),
        sem_remessa=remessa_dia.MOTIVO_FORA_SICOOB)
    assert texto == (f"APTO · {remessa_dia.MOTIVO_FORA_SICOOB} — "
                     f"{confirmacao.PAGUE_PELO_HTML}")
    assert estado == "ok"


def test_em_conta_de_outro_banco_o_reembolso_barrado_pela_identificacao_e_ambar():
    """O reembolso sem quem recebe nem sempre chega como `MOTIVO_REEMBOLSO`:
    quando a identificação achou um problema, o impedimento é o recado DELA
    (com o nome do aviso), copiado do `reembolso_impedimento` da linha. É a
    mesma família — não se sabe a quem pagar — e pinta igual."""
    recado = "o aviso manda pagar PESSOA DE EXEMPLO, que não está no cadastro"
    texto, estado = confirmacao.situacao_da_linha(
        {"status": "APTO* (reembolso)", "reembolso_impedimento": recado},
        _cand(impedimento=recado, reembolso=True),
        sem_remessa=remessa_dia.MOTIVO_FORA_SICOOB)
    assert texto == f"APTO* (reembolso) · {recado}"
    assert estado == "atencao"


def test_conta_sicoob_com_cadastro_incompleto_continua_ambar():
    for motivo in (remessa_dia.MOTIVO_SEM_CONVENIO, remessa_dia.MOTIVO_SEM_BANCO,
                   remessa_dia.MOTIVO_CONTA_DESCONHECIDA):
        assert confirmacao.situacao_da_linha(
            {"status": "APTO"}, _cand(), sem_remessa=motivo)[1] == "atencao"


def test_o_resumo_da_conta_de_outro_banco_tambem_e_neutro():
    fora = confirmacao.Grupo("CONTA INTER", [], [],
                             sem_remessa=remessa_dia.MOTIVO_FORA_SICOOB)
    incompleta = confirmacao.Grupo("CONTA SICOOB", [], [],
                                   sem_remessa=remessa_dia.MOTIVO_SEM_CONVENIO)
    assert "conta sem remessa" not in confirmacao.resumo_da_conta(fora)
    assert confirmacao.PAGUE_PELO_HTML in confirmacao.resumo_da_conta(fora)
    assert confirmacao.resumo_da_conta(incompleta).endswith(
        f"conta sem remessa: {remessa_dia.MOTIVO_SEM_CONVENIO}")


JA_SAIU = "já saiu na remessa nº 000007 de 10/09/2026"


def test_o_que_ja_saiu_em_remessa_pede_olhada():
    """O que faz a pessoa parar vem PRIMEIRO, como na conferência da remessa:
    marcar é o mesmo pagamento duas vezes."""
    texto, estado = confirmacao.situacao_da_linha(
        {"status": "APTO"}, _cand(ja_enviado=JA_SAIU))
    assert texto == f"APTO · {JA_SAIU} · vai na remessa"
    assert estado == "atencao"


def test_o_que_ja_saiu_aparece_em_conta_de_outro_banco():
    """Um boleto que já saiu na remessa nº 7 aparecia verde "pague pelo HTML":
    pagá-lo à mão é pagá-lo em dobro."""
    texto, estado = confirmacao.situacao_da_linha(
        {"status": "APTO"}, _cand(ja_enviado=JA_SAIU),
        sem_remessa=remessa_dia.MOTIVO_FORA_SICOOB)
    assert texto == (f"APTO · {JA_SAIU} · {remessa_dia.MOTIVO_FORA_SICOOB} — "
                     f"{confirmacao.PAGUE_PELO_HTML}")
    assert estado == "atencao"


def test_o_que_ja_saiu_aparece_em_conta_sem_remessa():
    texto, estado = confirmacao.situacao_da_linha(
        {"status": "APTO"}, _cand(ja_enviado=JA_SAIU),
        sem_remessa=remessa_dia.MOTIVO_SEM_CONVENIO)
    assert texto == (f"APTO · {JA_SAIU} · conta sem remessa: "
                     f"{remessa_dia.MOTIVO_SEM_CONVENIO}")
    assert estado == "atencao"


def test_o_que_ja_saiu_aparece_sem_o_cadastro_de_contas():
    texto, estado = confirmacao.situacao_da_linha(
        {"status": "APTO"}, _cand(ja_enviado=JA_SAIU), contas_conferidas=False)
    assert texto == f"APTO · {JA_SAIU} · {confirmacao.SEM_IMPEDIMENTO_NA_LINHA}"
    assert estado == "atencao"


@pytest.mark.parametrize("sem_remessa", [
    "", remessa_dia.MOTIVO_FORA_SICOOB, remessa_dia.MOTIVO_SEM_CONVENIO])
def test_sem_saber_se_ja_saiu_a_linha_diz_e_pinta(sem_remessa):
    """Registro não consultado: "não saiu" seria palpite. A linha diz que não
    conferiu, em qualquer conta, e fica âmbar."""
    texto, estado = confirmacao.situacao_da_linha(
        {"status": "APTO"}, _cand(), sem_remessa=sem_remessa,
        envio_conferido=False)
    assert texto.startswith(f"APTO · {confirmacao.NAO_CONFERI_ENVIO} · ")
    assert estado == "atencao"


def test_reembolso_pinta_ambar_em_qualquer_conta():
    """O dinheiro vai para quem NÃO é o favorecido do lançamento — na
    conferência da remessa essa linha já nasce desmarcada."""
    for sem_remessa in ("", remessa_dia.MOTIVO_FORA_SICOOB):
        _texto, estado = confirmacao.situacao_da_linha(
            {"status": "APTO* (reembolso)", "reembolso": True},
            _cand(status="APTO* (reembolso)", reembolso=True),
            sem_remessa=sem_remessa)
        assert estado == "atencao"


def test_sem_o_cadastro_de_contas_nao_se_promete_remessa():
    texto, estado = confirmacao.situacao_da_linha(
        {"status": "APTO"}, _cand(), contas_conferidas=False)
    assert confirmacao.VAI_NA_REMESSA not in texto
    assert texto.startswith("APTO · ")
    assert estado == "ok"


def test_sem_analise_fica_so_o_veredito_da_planilha():
    assert confirmacao.situacao_da_linha({"status": "APTO"}, None) \
        == ("APTO", "ok")


def test_por_onde_e_o_dado_que_a_planilha_vai_ter():
    digitos = "".join(ch for ch in LINHA_BANCARIA if ch.isdigit())
    assert confirmacao.por_onde("Pix", "fornecedor@exemplo.com") \
        == "PIX  fornecedor@exemplo.com"
    assert confirmacao.por_onde("Boleto", digitos) == f"BOLETO  {LINHA_BANCARIA}"
    assert confirmacao.por_onde("Boleto", "") \
        == f"BOLETO  {remessa_dia.MOTIVO_SEM_CHAVE}"
    assert confirmacao.por_onde(
        "Boleto", "86860000026-5 70860161209-4 22026081001-8 61001177300-1"
    ).startswith("ARRECADAÇÃO  ")


# ==========================================================================
# As linhas da janela: por conta, ENTRAM e NÃO APTOS
# ==========================================================================
def _reg(ident, favorecido="Fornecedor Modelo Ltda", **mudancas):
    base = {"id": ident, "tipo": "Pix", "dados": "fornecedor@exemplo.com",
            "valor": 100.0, "descricao": "OBRA MODELO NF 1234",
            "favorecido": favorecido, "status": "APTO",
            "conferencia": "NF ✓", "obs": "", "oc": "1234",
            "centro_custo": "OBRA MODELO"}
    base.update(mudancas)
    return base


def _omit(ident, conta="CONTA A", motivo=regras.MOTIVO_SEM_PAGAR, **mudancas):
    base = {"conta": conta, "tipo": "Pix", "valor": 250.0,
            "descricao": "Locacao de equipamento", "motivo": motivo,
            "favorecido": "Locadora Modelo Ltda", "id": ident, "oc": "",
            "centro_custo": "OBRA MODELO", "dados": "",
            "obs": "Pix sem chave no cadastro — buscar no ERP",
            "conferencia": "(não cruzado)"}
    base.update(mudancas)
    return base


def test_cada_conta_mostra_o_que_entra_e_o_que_nao_entrou():
    resultado = relatorio.Resultado(
        contas={"CONTA B": [_reg("B1", "Zulu Materiais"),
                            _reg("B2", "Socio Ficticio"),
                            _reg("B3", "Alfa Materiais")],
                "CONTA A": [_reg("A1")]},
        omitidos=[_omit("A9"), _omit("B9", conta="CONTA B"),
                  _omit("X1", conta="EMPRESA - APENAS LANCAMENTO",
                        fora_do_recorte=True)])
    lancamentos = [_lanc("A1", plannedDate="2026-09-15")]
    grupos = confirmacao.grupos_da_confirmacao(
        resultado, None, lancamentos, destacar=["SOCIO FICTICIO"])

    assert [g.conta for g in grupos] == ["CONTA A", "CONTA B"], \
        "a conta ignorada não aparece: não é pendência a corrigir no ERP"
    assert [(ln.favorecido, ln.olhar) for ln in grupos[1].entram] == [
        ("Socio Ficticio", True), ("Alfa Materiais", False),
        ("Zulu Materiais", False)]
    entra, = grupos[0].entram
    assert entra.marcavel and entra.vencimento == _dt.date(2026, 9, 15)
    assert (entra.oc, entra.centro_custo) == ("1234", "OBRA MODELO")
    nao, = grupos[0].nao_aptos
    assert (nao.id, nao.situacao, nao.estado) \
        == ("A9", regras.MOTIVO_SEM_PAGAR, "erro")
    assert not nao.marcavel, "não apto não se força: corrige-se no ERP"
    assert nao.obs == "Pix sem chave no cadastro — buscar no ERP"


def test_o_nao_apto_traz_o_pagamento_do_cadastro_do_lancamento():
    """"Sem forma de pagar" só se corrige sabendo o que o cadastro TEM: a TED
    escrita à mão, a chave com um dígito a mais. O texto vem do lançamento do
    passo 1, porque a linha omitida só guarda o que se conseguiu apurar."""
    resultado = relatorio.Resultado({}, [_omit("A9", tipo="TED")])
    lancamentos = [_lanc("A9", tradePayablePaymentMethod="TED",
                         paidToBankAccount="  BANCO 001 AG 1234 CC 56789-0 ")]
    grupo, = confirmacao.grupos_da_confirmacao(resultado, None, lancamentos)
    nao, = grupo.nao_aptos
    assert nao.pagamento_no_cadastro == "BANCO 001 AG 1234 CC 56789-0"
    grupo, = confirmacao.grupos_da_confirmacao(resultado, None, [])
    assert grupo.nao_aptos[0].pagamento_no_cadastro == ""


def test_a_linha_que_entra_traz_a_situacao_da_remessa():
    resultado = relatorio.Resultado(
        {CONTA: [_reg("L1", tipo="Boleto", dados=LINHA_BANCARIA)]}, [])
    candidato = _cand("L1", conta_erp=CONTA,
                      impedimento=remessa_dia.MOTIVO_PARCIAL)
    analise = confirmacao.AnaliseRemessa(preparado={CONTA: [candidato]},
                                         sem_remessa={})
    grupo, = confirmacao.grupos_da_confirmacao(resultado, analise, [])
    linha, = grupo.entram
    assert linha.situacao == f"APTO · remessa não leva: {remessa_dia.MOTIVO_PARCIAL}"
    assert linha.estado == "atencao"
    assert linha.candidato is candidato
    assert linha.por_onde == f"BOLETO  {LINHA_BANCARIA}"


def test_a_conta_sem_remessa_aparece_no_grupo():
    resultado = relatorio.Resultado({CONTA: [_reg("L1")]}, [])
    analise = confirmacao.AnaliseRemessa(
        preparado={CONTA: [_cand("L1", conta_erp=CONTA)]},
        sem_remessa={CONTA: remessa_dia.MOTIVO_SEM_CONVENIO})
    grupo, = confirmacao.grupos_da_confirmacao(resultado, analise, [])
    assert grupo.sem_remessa == remessa_dia.MOTIVO_SEM_CONVENIO
    assert grupo.entram[0].estado == "atencao"


def test_ja_pago_nao_e_pergunta():
    """Desmarcar um já pago não faria nada — o `montar_registros` não omite
    linha paga —, e marca que não tem efeito é mentira na tela."""
    resultado = relatorio.Resultado(
        {"CONTA A": [_reg("A1"), _reg("A2", status="JÁ PAGO em 13/09/2026")]},
        [])
    grupo, = confirmacao.grupos_da_confirmacao(
        resultado, None, [_lanc("A1"), _lanc("A2", paid=True)])
    assert [ln.id for ln in grupo.entram] == ["A1"]


def test_no_reembolso_quem_recebe_e_a_pessoa_do_aviso():
    """QUEM RECEBE dizia o fornecedor do lançamento, e o dinheiro vai para
    outra pessoa: a coluna mostra quem recebe de verdade e de que compra é o
    reembolso, e a linha fica âmbar."""
    resultado = relatorio.Resultado({"CONTA A": [
        _reg("R1", status="APTO* (reembolso)", reembolso=True,
             reembolso_nome="PESSOA DE EXEMPLO"),
        _reg("R2", status="APTO* (reembolso)", reembolso=True,
             reembolso_nome=""),
        _reg("N1")]}, [])
    grupo, = confirmacao.grupos_da_confirmacao(resultado, None, [])
    por_id = {ln.id: ln for ln in grupo.entram}
    assert por_id["R1"].quem_recebe \
        == "PESSOA DE EXEMPLO (reembolso de Fornecedor Modelo Ltda)"
    assert por_id["R2"].quem_recebe == "? (reembolso de Fornecedor Modelo Ltda)"
    assert por_id["R1"].estado == por_id["R2"].estado == "atencao"
    assert por_id["N1"].quem_recebe == "Fornecedor Modelo Ltda"
    assert por_id["N1"].estado == "ok"


def test_o_que_ja_saiu_em_remessa_nasce_desmarcado():
    """Como na conferência da remessa: marcar é o mesmo pagamento duas vezes,
    e isso exige um clique de quem leu o aviso."""
    resultado = relatorio.Resultado(
        {CONTA: [_reg("L1"), _reg("L2")]}, [_omit("L9", conta=CONTA)])
    analise = confirmacao.AnaliseRemessa(preparado={CONTA: [
        _cand("L1", conta_erp=CONTA),
        _cand("L2", conta_erp=CONTA, ja_enviado=JA_SAIU)]})
    grupo, = confirmacao.grupos_da_confirmacao(resultado, analise, [])
    por_id = {ln.id: ln for ln in grupo.entram + grupo.nao_aptos}
    assert confirmacao.marcada_de_inicio(por_id["L1"]) is True
    assert confirmacao.marcada_de_inicio(por_id["L2"]) is False
    assert confirmacao.marcada_de_inicio(por_id["L9"]) is False


#: Cadastro de mentira, no formato de `carregar_fornecedores`.
_MARCADA = {"CONCESSIONARIA LUZ": {"so_marcador": True}}


def test_o_marcador_de_recorrencia_nao_ocupa_a_janela():
    """A queixa de 20/08/2026: três linhas de R$ 1,00 todo dia, sabidas.

    `so_marcador` é decisão tomada no cadastro; repeti-la em vermelho todo
    dia ensina a pular os vermelhos — e é neles que mora o pagamento que
    ficou de fora. A marca é sobre o R$ 1,00 DESTE fornecedor: a conta de
    verdade dele que não entrou aparece, e o R$ 1,00 de outro também."""
    marcador = _omit("M1", motivo=regras.MOTIVO_SIMBOLICO, valor=1.0,
                     favorecido="CONCESSIONARIA LUZ S/A")
    outro = _omit("M2", motivo=regras.MOTIVO_SIMBOLICO, valor=1.0,
                  favorecido="Fornecedor Modelo Ltda")
    de_verdade = _omit("M3", valor=87.50, favorecido="CONCESSIONARIA LUZ S/A")
    resultado = relatorio.Resultado({}, [marcador, outro, de_verdade])

    grupo, = confirmacao.grupos_da_confirmacao(resultado, None, [],
                                               fornecedores=_MARCADA)
    assert [ln.id for ln in grupo.nao_aptos] == ["M2", "M3"]
    grupo, = confirmacao.grupos_da_confirmacao(resultado, None, [])
    assert [ln.id for ln in grupo.nao_aptos] == ["M1", "M2", "M3"], \
        "sem cadastro, nada é escondido"


def test_rodape_e_o_que_volta_saem_das_mesmas_marcas():
    resultado = relatorio.Resultado(
        {"CONTA A": [_reg("A1", valor=100.0), _reg("A2", valor=250.5),
                     _reg("A3", valor=10.0)]},
        [_omit("A8"), _omit("A9")])
    grupos = confirmacao.grupos_da_confirmacao(resultado, None, [])
    por_id = {ln.id: ln for ln in grupos[0].entram}
    marcas = [(por_id["A1"], True), (por_id["A2"], False),
              (por_id["A3"], True)]
    assert confirmacao.resumo(grupos, marcas) == (2, 110.0, 1, 2)
    assert confirmacao.nao_confirmados(marcas) == {"A2"}
    assert confirmacao.nao_confirmados(
        [(ln, True) for ln, _m in marcas]) == set()


def test_a_frase_do_rodape_conta_os_nao_aptos():
    assert confirmacao.frase_do_rodape(2, 110.0, 1, 2) \
        == "2 marcados  ·  R$ 110,00  ·  1 fica de fora  ·  2 não aptos"
    assert confirmacao.frase_do_rodape(1, 10.0, 0, 1) \
        == "1 marcado  ·  R$ 10,00  ·  1 não apto"
    assert confirmacao.frase_do_rodape(3, 30.0, 0, 0) \
        == "3 marcados  ·  R$ 30,00"


def test_a_cor_da_linha_na_tela():
    """Desmarcado é vermelho seja qual for o dado; não apto é sempre vermelho;
    a linha marcada fica com o próprio estado — e a que entra SEM dado de
    pagamento é âmbar, não vermelha: ela entra na planilha, quem não a leva é
    a remessa."""
    resultado = relatorio.Resultado(
        {"CONTA A": [_reg("A1"),
                     _reg("A2", dados="", status="ATENÇÃO — sem dados de pgto")]},
        [_omit("A9")])
    grupo, = confirmacao.grupos_da_confirmacao(resultado, None, [])
    por_id = {ln.id: ln for ln in grupo.entram + grupo.nao_aptos}
    assert confirmacao.estado_na_tela(por_id["A1"], True) == "ok"
    assert confirmacao.estado_na_tela(por_id["A2"], True) == "atencao"
    assert confirmacao.estado_na_tela(por_id["A1"], False) == "erro"
    assert confirmacao.estado_na_tela(por_id["A9"], True) == "erro"
