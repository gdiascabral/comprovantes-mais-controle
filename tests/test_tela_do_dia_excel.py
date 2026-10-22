"""O Excel do dia a partir da tela. Modelo fictício que imita o real: a linha
da «Movimentações» sai da fórmula, não de um deslocamento fixo."""
from decimal import Decimal as D

import openpyxl
import pytest

from conciliacao.aportadores import ler_regras
from conciliacao.workbook import WorkbookError
from tela_do_dia import excel_do_dia as xd
from tests.modelo_ficticio_do_dia import HOJE, PLANILHA, ROTULOS, carregar_mapa

ALFA, BETA, GAMA = ROTULOS
SOCIO = "Socio Ficticio"


@pytest.fixture
def mapa(tmp_path):
    return carregar_mapa(tmp_path)


def _modelo(caminho, cabecalho_mov="VALOR\nDEFINIDO (R$)"):
    wb = openpyxl.Workbook()
    p = wb.active
    p.title = "Painel"
    mov = wb.create_sheet("Movimentações")
    reg = wb.create_sheet("Regras")
    p["H7"] = "APORTE\nDEFINIDO (R$)"
    for r, rotulo in zip((8, 9, 10), ROTULOS):
        p[f"B{r}"] = rotulo
        p[f"G{r}"] = f'=IF($C{r}>=$D{r}+$E{r},"—",$D{r}+$E{r}-$C{r})'
    mov["H7"] = cabecalho_mov
    for linha_mov, linha_regra in ((9, 5), (10, 6), (13, 7)):
        mov[f"B{linha_mov}"] = f'=IF(Regras!$C{linha_regra}="","",Regras!$C{linha_regra})'
        mov[f"I{linha_mov}"] = f'=IF($H{linha_mov}<>"",$H{linha_mov},"")'
    mov["B12"] = "2 · O QUE OS INVESTIDORES MANDAM"
    mov["B53"] = "3 · À MÃO"
    for linha in range(54, 74):
        mov[f"I{linha}"] = f'=IF($H{linha}="","",$H{linha})'
    for linha, (destino, aportador, painel) in zip((5, 6, 7), (
            (ALFA, BETA, "sim"), (GAMA, BETA, "sim"), (ALFA, SOCIO, "não"))):
        reg[f"B{linha}"], reg[f"C{linha}"] = destino, aportador
        reg[f"E{linha}"], reg[f"F{linha}"] = "Aporte", painel
    wb.save(caminho)
    return caminho


def _gerar(tmp_path, mapa, definidos, a_mao=(), cabecalho_mov="VALOR\nDEFINIDO (R$)", **k):
    modelo = _modelo(tmp_path / "MODELO.xlsx", cabecalho_mov=cabecalho_mov)
    destino = tmp_path / "saida" / "21 10 - completa.xlsx"
    xd.gerar(modelo=modelo, destino=destino, data=HOJE, mapping=mapa, planilha=PLANILHA,
             saldos={8: D("100.50"), 9: D("-20.00"), 10: None},
             a_pagar={8: (D("300.00"), 2)}, regras=ler_regras(modelo),
             definidos=definidos, a_mao=a_mao, **k)
    return openpyxl.load_workbook(destino)


def test_linhas_da_movimentacao_saem_da_formula(tmp_path):
    ws = openpyxl.load_workbook(_modelo(tmp_path / "M.xlsx"))["Movimentações"]
    assert xd.linhas_da_movimentacao(ws) == {5: 9, 6: 10, 7: 13}


def test_o_excel_do_dia_leva_os_numeros_da_tela(tmp_path, mapa):
    wb = _gerar(tmp_path, mapa, {(ALFA, SOCIO): D("199.50")})
    p, mov = wb["Painel"], wb["Movimentações"]
    assert p["E4"].value.date() == HOJE
    assert (p["C8"].value, p["D8"].value, p["J8"].value, p["K8"].value) == (100.5, 300, 2, None)
    assert (p["C9"].value, p["D9"].value, p["K9"].value) == (-20, 0, 0)
    assert p["C10"].value is None                       # sem saldo: vazio, nunca 0
    assert p["H8"].value == 199.5 and p["H9"].value is None
    assert mov["H13"].value == 199.5
    # Participante da MESMA conta sem valor leva ZERO: vazio faria a coluna I
    # cair no VALOR SUGERIDO e somar dinheiro que ninguém definiu.
    assert mov["H9"].value == 0
    assert mov["H10"].value is None                     # regra de outra conta
    assert p["G8"].value.startswith("=IF(")             # fórmula preservada


def test_so_numero_e_escrito(tmp_path, mapa):
    modelo = _modelo(tmp_path / "MODELO.xlsx")
    extras = xd.celulas_de_aporte(modelo, PLANILHA, mapa, ler_regras(modelo),
                                  {(ALFA, SOCIO): D("1.00")})
    # O dicionário EXATO: `all(...)` passaria com um dicionário vazio.
    assert extras == {"Painel": {"H8": 1.0}, "Movimentações": {"H9": 0.0, "H13": 1.0}}
    assert all(type(v) is float for celulas in extras.values() for v in celulas.values())


@pytest.mark.parametrize("definidos, trecho", [
    ({(GAMA, SOCIO): D("10")}, "Regras"),
    ({(ALFA, SOCIO): D("-1")}, "negativo"),
])
def test_recusa_o_que_nao_tem_lugar_no_modelo(tmp_path, mapa, definidos, trecho):
    with pytest.raises(WorkbookError, match=trecho):
        _gerar(tmp_path, mapa, definidos)


def test_cabecalho_mudado_para_tudo(tmp_path, mapa):
    with pytest.raises(WorkbookError, match="VALOR DEFINIDO"):
        _gerar(tmp_path, mapa, {(ALFA, SOCIO): D("1")}, cabecalho_mov="VALOR (R$)")


# --- movimentos à mão: seção 3 da Movimentações, linhas 54 a 73 ---

MAO = [(None, ALFA, D("50.00")), (ALFA, BETA, D("10.25"))]


def test_movimentos_a_mao_vao_para_b_c_h_da_secao_3(tmp_path, mapa):
    wb = _gerar(tmp_path, mapa, {}, a_mao=MAO)
    mov = wb["Movimentações"]
    assert [mov[f"{c}54"].value for c in "BCH"] == [None, ALFA, 50.0]
    assert [mov[f"{c}55"].value for c in "BCH"] == [ALFA, BETA, 10.25]
    assert [mov[f"{c}56"].value for c in "BCH"] == [None, None, None]
    assert mov["I54"].value == '=IF($H54="","",$H54)'   # fórmula do modelo intacta
    assert wb["Painel"]["H8"].value is None             # sem aporte definido


def test_a_mao_e_aporte_definido_juntos(tmp_path, mapa):
    wb = _gerar(tmp_path, mapa, {(ALFA, SOCIO): D("199.50")}, a_mao=MAO)
    assert wb["Movimentações"]["H13"].value == 199.5
    assert wb["Movimentações"]["H54"].value == 50.0


def test_a_mao_zero_ou_negativo_fica_de_fora_como_na_tela(tmp_path, mapa):
    wb = _gerar(tmp_path, mapa, {}, a_mao=[(None, ALFA, D("0")), (None, BETA, D("-5")),
                                           (ALFA, GAMA, D("7"))])
    mov = wb["Movimentações"]
    assert [mov[f"{c}54"].value for c in "BCH"] == [ALFA, GAMA, 7.0]
    assert mov["H55"].value is None


def test_vinte_movimentos_cabem(tmp_path, mapa):
    wb = _gerar(tmp_path, mapa, {}, a_mao=[(None, ALFA, D(i + 1)) for i in range(20)])
    assert wb["Movimentações"]["H73"].value == 20.0


def test_vinte_e_um_movimentos_sao_recusados(tmp_path, mapa):
    with pytest.raises(WorkbookError, match="20"):
        _gerar(tmp_path, mapa, {}, a_mao=[(None, ALFA, D(i + 1)) for i in range(21)])
    assert not (tmp_path / "saida").exists()


@pytest.mark.parametrize("origem, destino", [("=1+1", ALFA), (ALFA, "=SUM(A1)"), (ALFA, ""),
                                             (ALFA, None)])
def test_a_mao_com_formula_ou_sem_destino_e_recusada(tmp_path, mapa, origem, destino):
    with pytest.raises(WorkbookError):
        _gerar(tmp_path, mapa, {}, a_mao=[(origem, destino, D("1"))])
    assert not (tmp_path / "saida").exists()


def test_a_mao_para_destino_fora_do_painel_e_recusado(tmp_path, mapa):
    with pytest.raises(WorkbookError, match="Conta Fora.*não é conta do painel"):
        _gerar(tmp_path, mapa, {}, a_mao=[(ALFA, "Conta Fora", D("5"))])
    assert not (tmp_path / "saida").exists()


def test_a_mao_com_origem_de_fora_do_painel_continua_valendo(tmp_path, mapa):
    wb = _gerar(tmp_path, mapa, {}, a_mao=[("Socio de fora", ALFA, D("5"))])
    assert wb["Movimentações"]["B54"].value == "Socio de fora"


def test_par_repetido_nas_regras_e_recusado(tmp_path, mapa):
    modelo = _modelo(tmp_path / "MODELO.xlsx")
    regras = ler_regras(modelo)
    with pytest.raises(WorkbookError, match="mais de uma vez"):
        xd.celulas_de_aporte(modelo, PLANILHA, mapa, regras + [regras[2]],
                             {(ALFA, SOCIO): D("1")})


def test_a_mao_com_cabecalho_mudado_para_tudo(tmp_path, mapa):
    with pytest.raises(WorkbookError, match="VALOR DEFINIDO"):
        _gerar(tmp_path, mapa, {}, a_mao=MAO, cabecalho_mov="VALOR (R$)")


def test_a_mao_com_valor_que_nao_e_decimal_e_recusado(tmp_path, mapa):
    with pytest.raises(WorkbookError, match="Decimal"):
        _gerar(tmp_path, mapa, {}, a_mao=[(None, ALFA, 1.5)])


# --- leitura que não pode virar Excel ---

def test_leitura_interrompida_nunca_vira_excel(tmp_path, mapa):
    with pytest.raises(WorkbookError, match="interrompida"):
        _gerar(tmp_path, mapa, {}, leitura_interrompida=True)
    assert not (tmp_path / "saida").exists()


def test_leitura_ausente_nunca_vira_excel(tmp_path, mapa):
    modelo = _modelo(tmp_path / "MODELO.xlsx")
    destino = tmp_path / "saida" / "d.xlsx"
    with pytest.raises(WorkbookError, match="ausente"):
        xd.gerar(modelo=modelo, destino=destino, data=HOJE, mapping=mapa, planilha=PLANILHA,
                 saldos={8: D("1")}, a_pagar=None, regras=[], definidos={})
    assert not destino.exists()
