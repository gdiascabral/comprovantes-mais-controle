"""Células EXTRAS do painel do dia (aporte definido, valor por aportador).
Modelo fictício no tmp_path — o real fica fora do repositório."""
from decimal import Decimal

import openpyxl
import pytest

from conciliacao.models import RowFill
from conciliacao.workbook import WorkbookError, assert_untouched, build
from tests.modelo_ficticio_do_dia import HOJE, PLANILHA, ROTULOS, carregar_mapa


@pytest.fixture
def mapa(tmp_path):
    return carregar_mapa(tmp_path)


@pytest.fixture
def modelo(tmp_path):
    wb = openpyxl.Workbook()
    p = wb.active
    p.title = "Painel"
    mov = wb.create_sheet("Movimentações")
    wb.create_sheet("Regras")
    p["H7"] = "APORTE\nDEFINIDO (R$)"
    for r, rotulo in zip((8, 9, 10), ROTULOS):
        p[f"B{r}"] = rotulo
        p[f"I{r}"] = f"=$C{r}-$D{r}+$H{r}"
    p["C12"] = "=SUM(C8:C10)"
    mov["H7"] = "VALOR\nDEFINIDO (R$)"
    mov["B9"] = '=IF(Regras!$C5="","",Regras!$C5)'
    mov["I9"] = '=IF($H9<>"",$H9,"")'
    caminho = tmp_path / "MODELO.xlsx"
    wb.save(caminho)
    return caminho


def _fills():
    return [RowFill(8, Decimal("100.50"), Decimal("300.00"), 2, None),
            RowFill(9, Decimal("0"), Decimal("0"), 0, 0),
            RowFill(10, None, Decimal("0"), 0, 0)]


def test_sem_extras_continua_como_antes(modelo, mapa, tmp_path):
    assert build(modelo, tmp_path / "s.xlsx", HOJE, _fills(), mapa, PLANILHA).linhas_escritas == 3


def test_extras_entram_e_sao_conferidos(modelo, mapa, tmp_path):
    destino = tmp_path / "s.xlsx"
    build(modelo, destino, HOJE, _fills(), mapa, PLANILHA,
          extras={"Painel": {"H8": 199.5}, "Movimentações": {"H9": 150.0}})
    wb = openpyxl.load_workbook(destino)
    assert wb["Painel"]["H8"].value == 199.5
    assert wb["Movimentações"]["H9"].value == 150.0
    assert wb["Painel"]["I8"].value == "=$C8-$D8+$H8"
    assert wb["Movimentações"]["I9"].value == '=IF($H9<>"",$H9,"")'


@pytest.mark.parametrize("valor", ["=1+1", "199,50", Decimal("1"), None, True])
def test_extra_que_nao_e_numero_e_recusado_antes_de_copiar(modelo, mapa, tmp_path, valor):
    destino = tmp_path / "s.xlsx"
    with pytest.raises(WorkbookError):
        build(modelo, destino, HOJE, _fills(), mapa, PLANILHA, extras={"Painel": {"H8": valor}})
    assert not destino.exists()


def test_extra_em_aba_que_o_modelo_nao_tem_e_recusado(modelo, mapa, tmp_path):
    destino = tmp_path / "s.xlsx"
    with pytest.raises(WorkbookError, match="Fantasma"):
        build(modelo, destino, HOJE, _fills(), mapa, PLANILHA, extras={"Fantasma": {"A1": 1.0}})
    assert not destino.exists()


@pytest.mark.parametrize("coord", [
    "I8",      # coluna de fórmula
    "C8",      # coluna que o programa já escreve
    "E4",      # célula da data
    "H20",     # fora da área conferida (abaixo da linha de totais)
    "N8",      # fora da área conferida (depois da última coluna)
    "H7",      # já tem texto no modelo
    "C12",     # já tem fórmula no modelo
])
def test_extra_do_painel_fora_do_permitido_e_recusado(modelo, mapa, tmp_path, coord):
    destino = tmp_path / "s.xlsx"
    with pytest.raises(WorkbookError, match=coord):
        build(modelo, destino, HOJE, _fills(), mapa, PLANILHA, extras={"Painel": {coord: 1.0}})
    assert not destino.exists()


@pytest.mark.parametrize("coord", ["I9", "B9", "H7"])
def test_extra_em_celula_com_valor_de_qualquer_aba_e_recusado(modelo, mapa, tmp_path, coord):
    destino = tmp_path / "s.xlsx"
    with pytest.raises(WorkbookError, match=coord):
        build(modelo, destino, HOJE, _fills(), mapa, PLANILHA,
              extras={"Movimentações": {coord: 1.0}})
    assert not destino.exists()


def test_celula_nao_extra_alterada_depois_do_build_continua_acusada(modelo, mapa, tmp_path):
    destino = tmp_path / "s.xlsx"
    extras = {"Painel": {"H8": 199.5}, "Movimentações": {"H9": 150.0}}
    build(modelo, destino, HOJE, _fills(), mapa, PLANILHA, extras=extras)
    assert assert_untouched(modelo, destino, HOJE, _fills(), PLANILHA, extras=extras) > 0

    wb = openpyxl.load_workbook(destino)
    wb["Movimentações"]["B9"] = "adulterada"
    wb.save(destino)
    wb.close()
    with pytest.raises(WorkbookError, match="Movimentações!B9"):
        assert_untouched(modelo, destino, HOJE, _fills(), PLANILHA, extras=extras)


def test_extra_gravado_errado_e_acusado(modelo, mapa, tmp_path):
    destino = tmp_path / "s.xlsx"
    extras = {"Movimentações": {"H9": 150.0}}
    build(modelo, destino, HOJE, _fills(), mapa, PLANILHA, extras=extras)
    wb = openpyxl.load_workbook(destino)
    wb["Movimentações"]["H9"] = 151.0
    wb.save(destino)
    wb.close()
    with pytest.raises(WorkbookError, match="Movimentações!H9"):
        assert_untouched(modelo, destino, HOJE, _fills(), PLANILHA, extras=extras)


# --- seção 3 «À MÃO» da Movimentações: B/C (origem/destino) aceitam TEXTO ---

def test_texto_na_secao_a_mao_entra_e_e_conferido(modelo, mapa, tmp_path):
    destino = tmp_path / "s.xlsx"
    extras = {"Movimentações": {"B54": "Conta Alfa - Banco", "C54": "Conta Beta - Banco",
                                "H54": 50.0, "C73": "Conta Gama - Banco", "H73": 7.0}}
    build(modelo, destino, HOJE, _fills(), mapa, PLANILHA, extras=extras)
    mov = openpyxl.load_workbook(destino)["Movimentações"]
    assert (mov["B54"].value, mov["C54"].value, mov["H54"].value) == (
        "Conta Alfa - Banco", "Conta Beta - Banco", 50.0)
    assert mov["B73"].value is None and mov["C73"].value == "Conta Gama - Banco"
    assert assert_untouched(modelo, destino, HOJE, _fills(), PLANILHA, extras=extras) > 0
    mov_wb = openpyxl.load_workbook(destino)
    mov_wb["Movimentações"]["C54"] = "Conta Gama - Banco"
    mov_wb.save(destino)
    with pytest.raises(WorkbookError, match="Movimentações!C54"):
        assert_untouched(modelo, destino, HOJE, _fills(), PLANILHA, extras=extras)


@pytest.mark.parametrize("aba, coord", [
    ("Movimentações", "B53"), ("Movimentações", "C74"), ("Movimentações", "D54"),
    ("Movimentações", "H55"), ("Movimentações", "A54"), ("Painel", "H8"),
    ("Regras", "B54"),
])
def test_texto_fora_de_b_c_54_a_73_continua_recusado(modelo, mapa, tmp_path, aba, coord):
    destino = tmp_path / "s.xlsx"
    with pytest.raises(WorkbookError, match="so numero|texto"):
        build(modelo, destino, HOJE, _fills(), mapa, PLANILHA,
              extras={aba: {coord: "Conta Alfa - Banco"}})
    assert not destino.exists()


@pytest.mark.parametrize("texto", ["=1+1", "=Regras!C5", "", "   "])
def test_formula_ou_texto_vazio_na_secao_a_mao_e_recusado(modelo, mapa, tmp_path, texto):
    destino = tmp_path / "s.xlsx"
    with pytest.raises(WorkbookError, match="B54"):
        build(modelo, destino, HOJE, _fills(), mapa, PLANILHA,
              extras={"Movimentações": {"B54": texto}})
    assert not destino.exists()


def test_texto_na_secao_a_mao_em_celula_ocupada_e_recusado(modelo, mapa, tmp_path):
    wb = openpyxl.load_workbook(modelo)
    wb["Movimentações"]["B54"] = "ja tem"
    wb.save(modelo)
    with pytest.raises(WorkbookError, match="B54"):
        build(modelo, tmp_path / "s.xlsx", HOJE, _fills(), mapa, PLANILHA,
              extras={"Movimentações": {"B54": "Conta Alfa - Banco"}})
