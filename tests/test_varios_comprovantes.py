# -*- coding: utf-8 -*-
"""Anexar N comprovantes num pagamento só (aba "Anexar (novo)"). Sem rede: dublê da API."""
from pathlib import Path

from anexar import matcher, matcher_somado, varios_comprovantes as vc

A400 = "400,00 - OBRA MODELO OC 5555 - 22-10.pdf"
A600 = "600,00 - OBRA MODELO OC 5555 - 22-10.pdf"


class ApiFalsa:
    def __init__(self, respostas=None, some=False, nao_verifica=(), listagem_falha=False):
        self.anexos, self.enviados = {}, []
        self.respostas, self.some, self.nao_verifica = list(respostas or []), some, set(nao_verifica)
        self.listagem_falha = listagem_falha

    def verificar_anexos(self, ids, log=print, progresso=None, cancelar=None):
        return {i: (-1 if i in self.nao_verifica else len(self.anexos.get(i, []))) for i in ids}

    def _anexos_brutos(self, pid):
        if self.listagem_falha:
            return {"__erro": 500}
        return list(self.anexos.get(pid, []))

    def anexar_por_api(self, pid, caminho, *, log=None, dry_run=False, tag=None):
        caminho = Path(caminho)
        self.enviados.append(caminho.name)
        if dry_run:
            return "dry_run"
        r = self.respostas.pop(0) if self.respostas else "anexado"
        if r == "anexado" and not self.some:
            self.anexos.setdefault(pid, []).append(
                {"name": caminho.name, "sizeInBytes": caminho.stat().st_size})
        return r


def _arquivos(tmp_path, tamanhos=(10, 11)):
    caminhos = []
    for nome, n in zip((A600, A400), tamanhos):
        p = tmp_path / nome
        p.write_bytes(b"x" * n)
        caminhos.append(p)
    return caminhos


def test_anexa_os_dois_e_prova_pelo_nome(tmp_path):
    api = ApiFalsa()
    r = vc.anexar_varios(api, "P1", _arquivos(tmp_path))
    assert r.estado == "anexado" and r.ok and api.enviados == [A600, A400]
    assert r.entraram == [A600, A400] and r.faltaram == []


def test_tamanhos_iguais_nao_sobem_nada(tmp_path):
    api = ApiFalsa()
    r = vc.anexar_varios(api, "P1", _arquivos(tmp_path, (10, 10)))
    assert r.estado == "erro:tamanhos_iguais" and api.enviados == []


def test_pagamento_que_ja_tem_anexo_nao_recebe_mais(tmp_path):
    api = ApiFalsa()
    api.anexos["P1"] = [{"name": "outro.pdf", "sizeInBytes": 5}]
    r = vc.anexar_varios(api, "P1", _arquivos(tmp_path))
    assert r.estado == "erro:ja_tem_anexo" and api.enviados == []


def test_rodada_nova_depois_de_parcial_nao_repete_o_que_entrou(tmp_path):
    """O 1º subiu numa rodada anterior: a rodada nova vê o anexo e não sobe nada,
    e diz quais dos N já estão lá e quais faltam."""
    api = ApiFalsa()
    arquivos = _arquivos(tmp_path)
    api.anexos["P1"] = [{"name": A600, "sizeInBytes": 10}]
    r = vc.anexar_varios(api, "P1", arquivos)
    assert r.estado == "erro:ja_tem_anexo" and api.enviados == []
    assert r.entraram == [A600] and r.faltaram == [A400]


def test_listagem_que_falha_nao_vira_sem_anexo(tmp_path):
    api = ApiFalsa(listagem_falha=True)
    r = vc.anexar_varios(api, "P1", _arquivos(tmp_path))
    assert r.estado == "erro:listagem:500" and api.enviados == [] and not r.ok


def test_erro_no_segundo_para_sem_tentar_de_novo(tmp_path):
    api = ApiFalsa(respostas=["anexado", "erro:upload:500"])
    r = vc.anexar_varios(api, "P1", _arquivos(tmp_path))
    assert r.estado.startswith("erro:parcial:") and not r.ok
    assert api.enviados == [A600, A400]                   # nenhum POST repetido
    assert r.entraram == [A600] and r.faltaram == [A400]
    assert A600 in r.estado and A400 in r.estado            # o relato diz quais


def test_erro_no_primeiro_nao_tenta_o_segundo(tmp_path):
    api = ApiFalsa(respostas=["erro:batch:500"])
    r = vc.anexar_varios(api, "P1", _arquivos(tmp_path))
    assert r.estado.startswith("erro:parcial:0/2") and api.enviados == [A600]
    assert r.entraram == [] and r.faltaram == [A600, A400]


def test_sem_a_prova_pelo_nome_nao_e_anexado(tmp_path):
    api = ApiFalsa(some=True)
    r = vc.anexar_varios(api, "P1", _arquivos(tmp_path))
    assert r.estado == "erro:nao_confirmado:0/2" and not r.ok
    assert r.entraram == [] and r.faltaram == [A600, A400]


def test_prova_parcial_pelo_nome_diz_qual_falta(tmp_path):
    api = ApiFalsa()
    arquivos = _arquivos(tmp_path)
    original = api.anexar_por_api

    def some_o_segundo(pid, caminho, **kw):
        r = original(pid, caminho, **kw)
        if Path(caminho).name == A400:
            api.anexos[pid] = [i for i in api.anexos[pid] if i["name"] != A400]
        return r

    api.anexar_por_api = some_o_segundo
    r = vc.anexar_varios(api, "P1", arquivos)
    assert r.estado == "erro:nao_confirmado:1/2"
    assert r.entraram == [A600] and r.faltaram == [A400]


def test_simular_nao_sobe(tmp_path):
    api = ApiFalsa()
    r = vc.anexar_varios(api, "P1", _arquivos(tmp_path), dry_run=True)
    assert r.estado == "dry_run" and api.anexos == {}


def test_nunca_passa_pela_tela():
    texto = Path(vc.__file__).read_text(encoding="utf-8")
    assert "anexar_um" not in texto and ".anexar(" not in texto


def _pend(paid_id="P1", valor=100000):
    return {"paidId": paid_id, "launchId": "L-" + paid_id, "valor": valor, "valores": [valor],
            "doc": "", "desc": "COMPRA OC 5555", "works": ["OBRA MODELO"], "data": "2210",
            "conta": "CONTA EXEMPLO - SICOOB", "favorecido": "Fornecedor Modelo Ltda",
            "dataFull": "22/10/2026", "categoria": ""}


def _fase(tmp_path, api, pagos, nomes=(A400, A600)):
    _arquivos(tmp_path)
    return vc.fase_das_somas(api, pagos, tmp_path, simular=False,
                             carregar_pdfs=lambda pasta, log: [matcher.parse_pdf(n) for n in nomes],
                             log=lambda _m: None)


def test_fase_das_somas_anexa_quem_ficou_sem_par(tmp_path):
    pe = _pend()
    matcher.casar([pe], [])
    somas = _fase(tmp_path, ApiFalsa(), [pe])
    assert [p["paidId"] for p in somas.anexados] == ["P1"] and somas.ok == 1
    assert pe["resultado"] == "anexado" and pe["pdf"] == f"{A600} + {A400}"


def test_fase_das_somas_ignora_nao_verificado(tmp_path):
    pe = _pend()
    matcher.casar([pe], [])
    api = ApiFalsa(nao_verifica={"P1"})
    somas = _fase(tmp_path, api, [pe])
    assert somas.anexados == [] and somas.nao_verificados == 1 and api.enviados == []


def test_fase_das_somas_nao_mexe_em_quem_ja_tem_anexo(tmp_path):
    pe = _pend()
    matcher.casar([pe], [])
    api = ApiFalsa()
    api.anexos["P1"] = [{"name": A600, "sizeInBytes": 10}]
    somas = _fase(tmp_path, api, [pe])
    assert somas.anexados == [] and somas.com_anexo == 1 and api.enviados == []


def test_fase_das_somas_nao_reusa_pdf_da_fase_1(tmp_path):
    pe, outro = _pend(), _pend("P2", 70000)
    matcher.casar([pe], [])
    outro.update(status="CERTEZA", pdf=A400)
    somas = _fase(tmp_path, ApiFalsa(), [pe, outro])
    assert somas.anexados == []


def test_fase_das_somas_duvida_vem_com_o_motivo_e_nao_anexa(tmp_path):
    """Dois conjuntos possíveis: nada sobe, e a dúvida leva o `motivo_soma`."""
    pe = _pend()
    matcher.casar([pe], [])
    nomes = (A400, A600, "300,00 - OBRA MODELO OC 5555 - 22-10.pdf",
             "700,00 - OBRA MODELO OC 5555 - 22-10.pdf")
    api = ApiFalsa()
    somas = _fase(tmp_path, api, [pe], nomes=nomes)
    assert somas.anexados == [] and api.enviados == []
    assert somas.duvidas == [pe] and pe["status"] == "SEM PAR"
    assert "conjuntos" in pe["motivo_soma"]
    assert any("DÚVIDA" in linha and "conjuntos" in linha for linha in vc.resumo_das_somas(somas))


def test_fase_das_somas_le_o_motivo_soma_de_antes(tmp_path):
    """Quem já virou dúvida de soma (status SEM PAR, `motivo_soma` gravado) não
    some do relato nem é somado de novo: entra nas dúvidas."""
    pe = _pend()
    matcher.casar([pe], [])
    pe["motivo_soma"] = pe["motivo_sem_par"] = "2 conjuntos de comprovantes somam o valor — anexe à mão"
    api = ApiFalsa()
    somas = _fase(tmp_path, api, [pe])
    assert somas.anexados == [] and somas.duvidas == [pe] and api.enviados == []


def test_fase_das_somas_relata_parcial(tmp_path):
    pe = _pend()
    matcher.casar([pe], [])
    somas = _fase(tmp_path, ApiFalsa(respostas=["anexado", "erro:upload:500"]), [pe])
    assert somas.ok == 0 and pe["resultado"].startswith("erro:parcial:")
    assert pe["entraram"] == [A600] and pe["faltaram"] == [A400]
    texto = "\n".join(vc.resumo_das_somas(somas))
    assert f"entraram: {A600}" in texto and f"faltaram: {A400}" in texto


def test_anexado_registra_o_criterio_da_prova(tmp_path):
    r = vc.anexar_varios(ApiFalsa(), "P1", _arquivos(tmp_path))
    assert r.estado == "anexado" and r.provado_por == "nome"
    r2 = vc.anexar_varios(ApiFalsa(some=True), "P1", _arquivos(tmp_path))
    assert r2.provado_por == ""


def test_parcial_rele_a_listagem_e_separa_pelo_nome(tmp_path):
    """A API disse "anexado" no 1º, mas a releitura não o mostra pelo nome:
    ele NÃO conta como entrou. E o 2º, com erro que não é de batch, pode ter
    ficado lá."""
    api = ApiFalsa(respostas=["anexado", "erro:nao_confirmado"], some=True)
    r = vc.anexar_varios(api, "P1", _arquivos(tmp_path))
    assert r.estado.startswith("erro:parcial:")
    assert r.entraram == [] and r.faltaram == [A600, A400]
    assert "pode ter ficado lá" in r.estado


def test_parcial_por_erro_de_batch_nao_diz_que_pode_ter_ficado(tmp_path):
    api = ApiFalsa(respostas=["anexado", "erro:batch:500"])
    r = vc.anexar_varios(api, "P1", _arquivos(tmp_path))
    assert r.entraram == [A600] and r.faltaram == [A400]
    assert "pode ter ficado lá" not in r.estado


def test_soma_igual_a_pago_que_ja_tem_anexo_vira_duvida(tmp_path):
    """Rodada 1 anexou A600+A400 em P1. Rodada 2: P2, gêmeo (mesmo valor, data
    e obra) SEM PAR, não pode receber os mesmos dois de novo."""
    p1, p2 = _pend("P1"), _pend("P2")
    matcher.casar([p2], [])
    api = ApiFalsa()
    api.anexos["P1"] = [{"name": A600, "sizeInBytes": 10}, {"name": A400, "sizeInBytes": 11}]
    somas = _fase(tmp_path, api, [p1, p2])
    assert somas.anexados == [] and api.enviados == []
    assert somas.duvidas == [p2] and p2["status"] == "SEM PAR"
    assert "já tem anexo" in p2["motivo_soma"]


def test_soma_igual_a_pago_nao_verificado_tambem_vira_duvida(tmp_path):
    p1, p2 = _pend("P1"), _pend("P2")
    matcher.casar([p2], [])
    api = ApiFalsa(nao_verifica={"P1"})
    somas = _fase(tmp_path, api, [p1, p2])
    assert somas.anexados == [] and api.enviados == [] and somas.duvidas == [p2]


def test_pago_de_outra_obra_com_anexo_nao_atrapalha(tmp_path):
    p1, p2 = _pend("P1"), _pend("P2")
    p1["works"] = ["OUTRA OBRA"]
    matcher.casar([p2], [])
    api = ApiFalsa()
    api.anexos["P1"] = [{"name": "x.pdf", "sizeInBytes": 3}]
    somas = _fase(tmp_path, api, [p1, p2])
    assert [p["paidId"] for p in somas.anexados] == ["P2"] and somas.ok == 1


def test_motivo_soma_de_quem_nao_esta_sem_par_e_ignorado(tmp_path):
    pe = _pend()
    pe.update(status="CERTEZA", pdf="outro.pdf", motivo_soma="motivo velho")
    somas = _fase(tmp_path, ApiFalsa(), [pe])
    assert somas.duvidas == [] and somas.anexados == []


def test_motivo_soma_velho_nao_prende_quem_esta_sem_par_agora(tmp_path):
    """SEM PAR nesta rodada pelo matcher (motivo_sem_par = nenhum PDF) com um
    `motivo_soma` que sobrou de antes: é somado de novo."""
    pe = _pend()
    pe["motivo_soma"] = "motivo velho"
    matcher.casar([pe], [])
    somas = _fase(tmp_path, ApiFalsa(), [pe])
    assert [p["paidId"] for p in somas.anexados] == ["P1"] and somas.duvidas == []
    assert "motivo_soma" not in pe


def test_o_motivo_que_abre_a_porta_e_o_do_matcher_somado():
    assert vc.matcher_somado.MOTIVO_NENHUM_PDF is matcher_somado.MOTIVO_NENHUM_PDF


def test_juntar_fim_soma_as_duas_fases():
    somas = vc.Somas(anexados=[{"resultado": "anexado"}], duvidas=[{}])
    assert vc.juntar_fim((3, 4, 1, 5, "r.xlsx"), somas) == (4, 5, 2, 3, "r.xlsx")


def test_relatorio_ganha_as_abas_das_somas(tmp_path):
    from openpyxl import Workbook, load_workbook
    caminho = tmp_path / "relatorio.xlsx"
    Workbook().save(caminho)
    pe = dict(_pend(), pdfs=[A600, A400], motivo="soma exata", resultado="anexado")
    vc.acrescentar_ao_relatorio(str(caminho), vc.Somas(anexados=[pe]))
    assert "ANEXADOS POR SOMA" in load_workbook(caminho).sheetnames
