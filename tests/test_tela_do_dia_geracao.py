"""O Resultado da tela: liberados, editados e à mão — e a remessa sem os à mão."""
import unicodedata
from datetime import date
from types import SimpleNamespace

import pytest

from pagamentos_dia import confirmacao, html_pagamentos as hp, relatorio, remessa_dia
from pagamentos_dia import regras_pagamento as regras
from tela_do_dia import conferencia as cf, geracao as ge
from tela_do_dia import pagamentos as pg

DIA = date(2026, 10, 22)
CONTA = "CONTA EXEMPLO - SICOOB"
NF = lambda n: {"filename": n, "tagName": "Nota Fiscal", "extension": ".pdf",  # noqa: E731
                "downloadUrl": f"https://exemplo.invalid/{n}.pdf"}
#: CPF fictício (o mesmo de `tests/test_remessa_dia.py`): com ele, APTO nasce marcado.
CONTATOS = {"FORNECEDOR MODELO LTDA": "52998224725"}


def _lanc(ident, valor=100.0):
    return {"id": ident, "tradePayableId": f"T-{ident}", "paidTo": "Fornecedor Modelo Ltda",
            "remainingValue": valor, "plannedDate": "2026-10-22", "tradePayableAccount": {"name": CONTA},
            "tradePayablePaymentMethod": "Pix", "paidToBankAccount": "PIX EMAIL fulano@exemplo.com",
            "documentNumber": "1234", "costCentreDetails": [{"workName": "OBRA MODELO"}], "paid": False}


def _entradas(lancs, anexos):
    return confirmacao.Entradas(selecionados=lancs, anexos=anexos, overviews={}, textos={},
                                urls_ocr=set(), cadastro_reembolso=None, regras_fornecedor={},
                                participantes={}, periodo=(DIA, DIA))


@pytest.fixture
def dia():
    e = _entradas([_lanc("L1"), _lanc("L2"), _lanc("L3")], {"T-L1": [NF("n1")], "T-L3": [NF("n3")]})
    grupos = confirmacao.grupos_da_confirmacao(confirmacao.remontar(e), None, e.selecionados)
    return e, grupos


def _conf(eventos=(), a_mao=()):
    return cf.Conferencia.da_nuvem(list(eventos), list(a_mao))


A_MAO_L9 = {"id": 4, "conta": CONTA, "valor": "30.00", "descricao": "parcela combinada",
            "motivo": "acordo", "favorecido": "Fornecedor Modelo Ltda", "tipo": "Pix",
            "dados": "PIX EMAIL fulano@exemplo.com", "lancamento_id": "L9", "centro_custo": "OBRA MODELO",
            "nf": "5678", "valor_em_aberto": "80.00"}


def _html(r, ident):
    return next(x for c in hp.contas_do_html_geral(r) for x in c["entries"] if x["id"] == ident)


# ----------------------------------------------------------------- do plano
def test_sem_conferencia_e_o_mesmo_da_entrega_1(dia):
    e, grupos = dia
    assert ge.resultado_da_tela(e, grupos, _conf()) == confirmacao.remontar(e, pg.ids_que_nao_saem(grupos))


def test_liberado_entra_marcado_como_liberado_e_nao_como_apto(dia):
    e, grupos = dia
    r = ge.resultado_da_tela(e, grupos, _conf([{"id": 1, "alvo": "L2", "acao": cf.LIBERAR,
                                                "motivo": "nota chega amanhã", "quem_nome": "Fulano"}]))
    reg = ge.registro_atual(r, "L2")
    assert reg["status"].startswith("LIBERADO — " + regras.MOTIVO_SEM_ANEXO)
    assert "liberado na tela por Fulano" in reg["obs"]
    cand = remessa_dia.preparar(ge.para_a_remessa(r).contas)[CONTA]
    assert not next(c for c in cand if c.id == "L2").apto        # nasce desmarcado na Remessa


def test_desmarcado_vai_para_os_omitidos(dia):
    e, grupos = dia
    r = ge.resultado_da_tela(e, grupos, _conf([{"id": 1, "alvo": "L3", "acao": cf.DESMARCAR}]))
    assert any(o["id"] == "L3" and o["motivo"] == regras.MOTIVO_NAO_CONFIRMADO for o in r.omitidos)


def test_edicao_de_forma_troca_o_dado_e_avisa_corrigir_no_erp(dia):
    e, grupos = dia
    r = ge.resultado_da_tela(e, grupos, _conf([{
        "id": 1, "alvo": "L1", "acao": cf.EDITAR, "motivo": "chave errada no ERP",
        "campos": {"tipo": "Pix", "dados": "PIX EMAIL outro@exemplo.com"},
        "antes": {"tipo": "Pix", "dados": "fulano@exemplo.com"}, "corrigir_no_erp": True}]))
    reg = ge.registro_atual(r, "L1")
    assert reg["dados"] == "PIX EMAIL outro@exemplo.com" and reg["pix_qr"] == ""
    assert reg["status"] == "EDITADO — forma de pagar"
    assert "corrigir no ERP: lançamento L1" in reg["obs"]


def test_so_descricao_mantem_o_status_e_vai_ao_html(dia):
    e, grupos = dia
    r = ge.resultado_da_tela(e, grupos, _conf([{"id": 1, "alvo": "L1", "acao": cf.EDITAR,
                                                "motivo": "texto melhor", "campos": {"descricao": "OBRA MODELO PISO 1234"}}]))
    reg = ge.registro_atual(r, "L1")
    assert reg["status"].startswith("APTO")
    assert _html(r, "L1")["descricao"] == "OBRA MODELO PISO 1234"


def test_atencao_grave_nao_some_com_a_edicao():
    d = cf.dobrar([cf.Evento.de_linha({"id": 1, "alvo": "X", "acao": cf.EDITAR, "motivo": "nova chave",
                                       "campos": {"tipo": "Pix", "dados": "PIX EMAIL a@exemplo.com"}})])["X"]
    reg = ge.aplicar_decisao({"id": "X", "status": "ATENÇÃO — valor do boleto diverge", "obs": ""}, d,
                             liberavel=False)
    assert reg["status"].startswith("ATENÇÃO — valor do boleto diverge")


def test_a_mao_entra_no_html_e_na_planilha_mas_nao_na_remessa(dia):
    e, grupos = dia
    r = ge.resultado_da_tela(e, grupos, _conf(a_mao=[A_MAO_L9]))
    reg = next(x for x in r.contas[CONTA] if x.get("a_mao") == "mao:4")
    assert reg["id"] == "L9" and reg["status"] == "À MÃO — acordo" and reg["valor"] == 30.0
    assert "mao:4" in ge.chaves_que_saem(r)
    assert all(not x.get("a_mao") for x in ge.para_a_remessa(r).contas[CONTA])
    # Achado 7 do preflight: no HTML o id do à mão é "<lançamento>#mao:<n>"; o link é o do lançamento.
    entrada = _html(r, "L9#mao:4")
    assert entrada["link"] == hp.URL_LANCAMENTO + "L9"
    assert entrada["descricao"].startswith("OBRA MODELO") and entrada["descricao"].endswith("5678")
    assert entrada["bloqueio"] == ""


# ------------------------------------------------- pontos de dinheiro (a mais)
def test_a_mao_com_link_e_o_lancamento_do_dia_nao_colidem_no_html(dia):
    """Preflight, achado 7: o mesmo id dava a mesma chave de "já paguei" e de "Conferido"."""
    e, grupos = dia
    r = ge.resultado_da_tela(e, grupos, _conf(a_mao=[{**A_MAO_L9, "lancamento_id": "L1"}]))
    ids = [x["id"] for c in hp.contas_do_html_geral(r) for x in c["entries"]]
    assert len(ids) == len(set(ids))
    assert "L1" in ids and "L1#mao:4" in ids
    assert _html(r, "L1#mao:4")["link"] == _html(r, "L1")["link"]
    assert ge.registro_atual(r, "L1").get("a_mao") is None     # o do dia, não o à mão


def test_a_mao_sem_link_tem_id_proprio_no_html_e_a_descricao_como_escrita(dia):
    e, grupos = dia
    sem_link = {k: v for k, v in A_MAO_L9.items()
                if k not in ("lancamento_id", "centro_custo", "nf", "valor_em_aberto")}
    r = ge.resultado_da_tela(e, grupos, _conf(a_mao=[sem_link]))
    entrada = _html(r, "mao:4")
    assert entrada["link"] == "" and entrada["descricao"] == "parcela combinada"
    reg = next(x for x in r.contas[CONTA] if x.get("a_mao") == "mao:4")
    assert "o Anexar não tem o que casar" in reg["obs"]


def test_a_mao_removido_nao_sai(dia):
    e, grupos = dia
    r = ge.resultado_da_tela(e, grupos, _conf([{"id": 9, "alvo": "mao:4", "acao": cf.REMOVER_A_MAO,
                                                "motivo": "lançado errado"}], a_mao=[A_MAO_L9]))
    assert "mao:4" not in ge.chaves_que_saem(r)


def test_forma_editada_sai_igual_nos_tres_geradores(dia, tmp_path):
    """Planilha, HTML e remessa leem o MESMO registro: nenhum sai com a chave antiga."""
    e, grupos = dia
    r = ge.resultado_da_tela(e, grupos, _conf([{
        "id": 1, "alvo": "L1", "acao": cf.EDITAR, "motivo": "chave errada no ERP",
        "campos": {"tipo": "Pix", "dados": "PIX EMAIL outro@exemplo.com"}}]))
    assert _html(r, "L1")["dados_original"] == "PIX EMAIL outro@exemplo.com"
    cand = next(c for c in remessa_dia.preparar(ge.para_a_remessa(r).contas, participantes=CONTATOS)[CONTA]
                if c.id == "L1")
    assert cand.chave == "PIX EMAIL outro@exemplo.com" and not cand.marcado
    from openpyxl import load_workbook
    caminho = relatorio.gerar_excel(r, tmp_path / "dia.xlsx", log=lambda *_: None)
    linhas = [[c.value for c in linha] for ws in load_workbook(caminho).worksheets
              for linha in ws.iter_rows()]
    editada = [ln for ln in linhas if ln and ln[0] == "Pix" and ln[5] == "EDITADO — forma de pagar"]
    assert [ln[1] for ln in editada] == ["PIX EMAIL outro@exemplo.com"]          # a planilha
    outras = [ln for ln in linhas if ln and ln[0] == "Pix" and ln not in editada]
    assert len(outras) == 1 and "outro@" not in str(outras[0][1])               # a L3, intacta


def test_forma_editada_limpa_o_que_era_da_cobranca_antiga():
    d = cf.dobrar([cf.Evento.de_linha({"id": 1, "alvo": "X", "acao": cf.EDITAR, "motivo": "boleto novo",
                                       "campos": {"tipo": "Pix", "dados": "PIX EMAIL a@exemplo.com"}})])["X"]
    reg = ge.aplicar_decisao({"id": "X", "status": "APTO", "obs": "", "tipo": "Boleto", "dados": "1" * 47,
                              "pix_qr": "000201abc", "linha_da_guia": "2" * 48, "parcial": True,
                              "conferencia": "NF 1234 × boleto: ok"}, d, liberavel=False)
    assert (reg["pix_qr"], reg["linha_da_guia"], reg["parcial"]) == ("", "", False)
    assert reg["conferencia"] == ""
    html = hp.contas_do_html_geral(relatorio.Resultado({CONTA: [{**reg, "valor": 10.0}]}, []))
    assert "parcial" not in html[0]["entries"][0]["bloqueio"]


def test_edicao_nao_tira_ja_pago():
    d = cf.dobrar([cf.Evento.de_linha({"id": 1, "alvo": "X", "acao": cf.EDITAR, "motivo": "nova chave",
                                       "campos": {"tipo": "Pix", "dados": "PIX EMAIL a@exemplo.com"}})])["X"]
    reg = ge.aplicar_decisao({"id": "X", "status": "JÁ PAGO em 21/10/2026", "obs": ""}, d, liberavel=True)
    assert reg["status"].startswith("JÁ PAGO")


def test_edicao_de_forma_nao_cede_o_sem_anexo():
    """Trocar a chave não traz o anexo: só a LIBERAÇÃO cede o "sem anexo"."""
    d = cf.dobrar([cf.Evento.de_linha({"id": 1, "alvo": "X", "acao": cf.EDITAR, "motivo": "nova chave",
                                       "campos": {"tipo": "Pix", "dados": "PIX EMAIL a@exemplo.com"}})])["X"]
    reg = ge.aplicar_decisao({"id": "X", "status": "ATENÇÃO — sem anexo", "obs": ""}, d, liberavel=False)
    assert reg["status"].startswith("ATENÇÃO — sem anexo")
    sem_dados = ge.aplicar_decisao({"id": "X", "status": "ATENÇÃO — sem dados de pgto", "obs": ""}, d,
                                   liberavel=False)
    assert sem_dados["status"] == "EDITADO — forma de pagar"


def test_liberado_que_seria_apto_vira_liberado_copiavel_e_desmarcado():
    """O gancho põe `STATUS_LIBERADO` (ATENÇÃO); a tela troca por LIBERADO: o HTML
    deixa copiar (o dono liberou com motivo) e a remessa continua sem marcar."""
    e = _entradas([_lanc("L5", 1.0)], {"T-L5": [NF("n5")]})
    grupos = confirmacao.grupos_da_confirmacao(confirmacao.remontar(e), None, e.selecionados)
    assert "L5" in ge.ids_nao_aptos(grupos)
    r = ge.resultado_da_tela(e, grupos, _conf([{"id": 1, "alvo": "L5", "acao": cf.LIBERAR,
                                                "motivo": "parcela simbólica combinada"}]))
    reg = ge.registro_atual(r, "L5")
    assert reg["status"].startswith("LIBERADO — " + regras.MOTIVO_SIMBOLICO)
    assert _html(r, "L5")["bloqueio"] == ""
    cand = remessa_dia.preparar(ge.para_a_remessa(r).contas, participantes=CONTATOS)[CONTA][0]
    assert not cand.apto and not cand.marcado


def test_descricao_editada_e_a_mesma_no_html_e_no_banco(dia):
    e, grupos = dia
    r = ge.resultado_da_tela(e, grupos, _conf([{"id": 1, "alvo": "L1", "acao": cf.EDITAR,
                                                "motivo": "texto melhor", "campos": {"descricao": "OBRA MODELO PISO 1234"}}]))
    cand = next(c for c in remessa_dia.preparar(ge.para_a_remessa(r).contas)[CONTA] if c.id == "L1")
    assert remessa_dia.etiqueta_do_banco(cand) == _html(r, "L1")["descricao"] == "OBRA MODELO PISO 1234"
    assert ge.registro_atual(r, "L1")["descricao"] == "OBRA MODELO PISO 1234"   # a planilha


def test_motivo_escrito_nao_vira_ordem_de_pagar_a_mao():
    """O motivo vai na Obs; "pagar à mão" escrito ali não pode virar o gatilho do HTML/remessa."""
    d = cf.dobrar([cf.Evento.de_linha({"id": 1, "alvo": "X", "acao": cf.EDITAR,
                                       "motivo": "senão vou pagar à mão", "campos": {"descricao": "OBRA MODELO 1"}})])["X"]
    reg = ge.aplicar_decisao({"id": "X", "status": "APTO", "obs": ""}, d, liberavel=False)
    assert "PAGAR A MAO" not in hp._norm(reg["obs"]) and "PAGAR À MÃO" not in reg["obs"]


def test_html_sem_a_mao_e_o_de_sempre(dia):
    e, grupos = dia
    r = confirmacao.remontar(e)
    assert [x["id"] for c in hp.contas_do_html_geral(r) for x in c["entries"]] == [
        str(x["id"]) for regs in r.contas.values() for x in regs]


# ------------------------------------------------------------- conserto (T5)
def _linha_ja_enviada(ident):
    return confirmacao.Linha(
        secao=confirmacao.ENTRA, id=ident, conta=CONTA, valor=100.0, favorecido="Fornecedor Modelo Ltda",
        tipo="Pix", dados="PIX EMAIL fulano@exemplo.com", por_onde="", vencimento=DIA, oc="",
        centro_custo="OBRA MODELO", descricao="", situacao="", estado="atencao", olhar=False,
        candidato=SimpleNamespace(ja_enviado="já saiu na remessa nº 7"))


def test_liberar_um_ja_enviado():
    e = _entradas([_lanc("L1")], {"T-L1": [NF("n1")]})
    grupos = [confirmacao.Grupo(CONTA, entram=[_linha_ja_enviada("L1")], nao_aptos=[])]
    assert pg.ids_que_nao_saem(grupos) == {"L1"}
    sem = ge.resultado_da_tela(e, grupos, _conf())
    assert ge.registro_atual(sem, "L1") not in [x for regs in sem.contas.values() for x in regs]
    assert "L1" not in ge.chaves_que_saem(sem)
    r = ge.resultado_da_tela(e, grupos, _conf([{"id": 1, "alvo": "L1", "acao": cf.LIBERAR,
                                                "motivo": "o banco recusou o arquivo"}]))
    reg = ge.registro_atual(r, "L1")
    assert reg in r.contas[CONTA]
    assert reg["status"].startswith("LIBERADO — já saiu na remessa nº 7")
    cand = remessa_dia.preparar(ge.para_a_remessa(r).contas, participantes=CONTATOS)[CONTA][0]
    assert cand.id == "L1" and not cand.apto and not cand.marcado
    assert _html(r, "L1")["bloqueio"] == ""


@pytest.mark.parametrize("motivo", [
    "senão vou pagar á mão", "PAGAR Á MÃO amanhã", "Pagar  a  mao",
    unicodedata.normalize("NFD", "senão vou pagar à mão")])
def test_gatilho_de_pagar_a_mao_em_qualquer_grafia(motivo):
    d = cf.dobrar([cf.Evento.de_linha({"id": 1, "alvo": "X", "acao": cf.EDITAR,
                                       "motivo": motivo, "campos": {"descricao": "OBRA MODELO 1"}})])["X"]
    reg = ge.aplicar_decisao({"id": "X", "status": "APTO", "obs": "", "tipo": "Pix",
                              "dados": "PIX EMAIL a@exemplo.com", "valor": 10.0}, d, liberavel=False)
    assert "PAGAR A MAO" not in hp._norm(reg["obs"]) and "PAGAR À MÃO" not in reg["obs"]
    assert hp.contas_do_html_geral(relatorio.Resultado({CONTA: [reg]}, []))[0]["entries"][0]["bloqueio"] == ""
    assert remessa_dia.preparar({CONTA: [reg]})[CONTA][0].impedimento != remessa_dia.MOTIVO_MAO


def test_motivo_sem_gatilho_fica_como_foi_escrito():
    d = cf.dobrar([cf.Evento.de_linha({"id": 1, "alvo": "X", "acao": cf.EDITAR,
                                       "motivo": "descrição não batia", "campos": {"descricao": "OBRA MODELO 1"}})])["X"]
    reg = ge.aplicar_decisao({"id": "X", "status": "APTO", "obs": ""}, d, liberavel=False)
    assert "(descrição não batia)" in reg["obs"]


def test_edicao_de_reembolso_vira_o_registro_de_reembolso(dia):
    e, grupos = dia
    r = ge.resultado_da_tela(e, grupos, _conf([{
        "id": 1, "alvo": "L1", "acao": cf.EDITAR, "motivo": "reembolso combinado",
        "campos": {"reembolso_nome": "Fulano Ficticio", "reembolso_documento": "52998224725"}}]))
    reg = ge.registro_atual(r, "L1")
    assert reg["reembolso"] is True
    assert (reg["reembolso_nome"], reg["reembolso_para"], reg["reembolso_documento"]) == (
        "Fulano Ficticio", "Fulano Ficticio", "52998224725")
    assert (reg["reembolso_origem"], reg["reembolso_impedimento"]) == ("tela", "")
    assert reg["status"] == "EDITADO — reembolso"
    assert reg["favorecido"] == "Fornecedor Modelo Ltda"                 # a planilha não troca
    assert _html(r, "L1")["favorecido"] == "Fulano Ficticio (reembolso de Fornecedor Modelo Ltda)"
    cand = next(c for c in remessa_dia.preparar(ge.para_a_remessa(r).contas, participantes=CONTATOS)[CONTA]
                if c.id == "L1")
    assert cand.reembolso and cand.favorecido == "Fulano Ficticio" and not cand.marcado


# ------------------------------- chave Pix editada à mão (pedido do dono, 22/10/2026)
#: CPF fictício válido, diferente do de CONTATOS: a chave que o dono digitou.
CPF_DA_CHAVE = "11144477735"


def _editar_l1(dados, tipo="Pix"):
    return _conf([{"id": 1, "alvo": "L1", "acao": cf.EDITAR, "motivo": "chave nova",
                   "campos": {"tipo": tipo, "dados": dados}}])


def _entradas_com_contatos():
    e = _entradas([_lanc("L1"), _lanc("L2"), _lanc("L3")], {"T-L1": [NF("n1")], "T-L3": [NF("n3")]})
    e = confirmacao.Entradas(**{**e.__dict__, "participantes": {"FORNECEDOR MODELO LTDA": "52998224725"}})
    return e, confirmacao.grupos_da_confirmacao(confirmacao.remontar(e), None, e.selecionados)


def test_mascara_do_documento():
    assert cf.mascarar_documento("52998224725") == "***.982.247-**"
    assert cf.mascarar_documento("11.222.333/0001-81") == "**.222.333/0001-**"
    assert cf.mascarar_documento("") == ""


def test_nota_da_chave_editada():
    assert cf.nota_da_chave_editada("PIX CPF 111.444.777-35", "52998224725") == (
        "chave Pix editada: CPF/CNPJ diferente do favorecido do ERP (***.982.247-**)")
    assert cf.nota_da_chave_editada("PIX CPF 529.982.247-25", "52998224725") == cf.NOTA_CHAVE_EDITADA
    assert cf.nota_da_chave_editada("PIX EMAIL a@exemplo.com", "52998224725") == cf.NOTA_CHAVE_EDITADA
    assert cf.nota_da_chave_editada("PIX CPF 111.444.777-35", "") == cf.NOTA_CHAVE_EDITADA


def test_chave_cpf_editada_vai_como_digitada_com_observacao_nos_tres_geradores(tmp_path):
    e, grupos = _entradas_com_contatos()
    r = ge.resultado_da_tela(e, grupos, _editar_l1("PIX CPF 111.444.777-35"))
    reg = ge.registro_atual(r, "L1")
    nota = "chave Pix editada: CPF/CNPJ diferente do favorecido do ERP (***.982.247-**)"
    assert reg["dados"] == "PIX CPF 111.444.777-35" and nota in reg["obs"]
    assert reg[remessa_dia.CHAVE_EDITADA] is True
    assert "52998224725" not in reg["obs"]                       # o documento do ERP vai mascarado
    assert nota in _html(r, "L1")["obs"]
    cand = next(c for c in remessa_dia.preparar(ge.para_a_remessa(r).contas,
                                                participantes=e.participantes)[CONTA] if c.id == "L1")
    assert cand.chave == "PIX CPF 111.444.777-35" and cand.documento_favorecido == CPF_DA_CHAVE
    from openpyxl import load_workbook
    caminho = relatorio.gerar_excel(r, tmp_path / "dia.xlsx", log=lambda *_: None)
    textos = [str(c.value) for ws in load_workbook(caminho).worksheets
              for linha in ws.iter_rows() for c in linha if c.value]
    assert any(nota in t for t in textos)


def test_chave_email_editada_leva_a_observacao_simples(dia):
    e, grupos = dia
    r = ge.resultado_da_tela(e, grupos, _editar_l1("PIX EMAIL outro@exemplo.com"))
    reg = ge.registro_atual(r, "L1")
    assert cf.NOTA_CHAVE_EDITADA in reg["obs"] and "diferente" not in reg["obs"]


def test_sem_edicao_nao_ha_marca_de_chave_editada(dia):
    e, grupos = dia
    r = ge.resultado_da_tela(e, grupos, _conf())
    assert not any(x.get(remessa_dia.CHAVE_EDITADA) for regs in r.contas.values() for x in regs)


# ------------------------- linha digitável editada: o valor confere (22/10/2026)


def _linha_bancaria(centavos: int) -> str:
    """Linha bancária fictícia com o valor pedido e os DVs fechando."""
    from pagamentos_dia import ocr_boleto
    d = ocr_boleto.digitos("34191.57007 00024.924375 24177.010006 9 15340000115000")
    d = d[:37] + f"{centavos:010d}"
    barra43 = d[0:4] + d[33:47] + d[4:9] + d[10:20] + d[21:31]
    return d[:32] + str(ocr_boleto._mod11_geral(barra43)) + d[33:]


def test_boleto_editado_que_bate_nao_anota(dia):
    e, grupos = dia
    r = ge.resultado_da_tela(e, grupos, _editar_l1(_linha_bancaria(10000), "Boleto"))
    reg = ge.registro_atual(r, "L1")
    assert reg["status"] == "EDITADO — forma de pagar"
    assert "valor da linha" not in reg["obs"] and "sem valor" not in reg["obs"]


def test_boleto_editado_sem_valor_entra_com_observacao(dia):
    e, grupos = dia
    r = ge.resultado_da_tela(e, grupos, _editar_l1(_linha_bancaria(0), "Boleto"))
    reg = ge.registro_atual(r, "L1")
    assert reg["status"] == "EDITADO — forma de pagar"
    assert "linha sem valor embutido — confira" in reg["obs"]


def test_boleto_editado_que_diverge_fica_em_atencao_com_a_observacao(dia):
    e, grupos = dia
    r = ge.resultado_da_tela(e, grupos, _editar_l1(_linha_bancaria(10050), "Boleto"))
    reg = ge.registro_atual(r, "L1")
    assert reg["status"].startswith("ATENÇÃO — valor da linha diverge")
    assert "EDITADO — forma de pagar" in reg["status"]
    assert "valor da linha R$ 100,50 ≠ pagamento R$ 100,00" in reg["obs"]
    assert _html(r, "L1")["bloqueio"]                         # o Copiar não vai às cegas
    cand = next(c for c in remessa_dia.preparar(ge.para_a_remessa(r).contas, participantes=CONTATOS)[CONTA]
                if c.id == "L1")
    assert not cand.marcado                                   # nasce desmarcado na Remessa
