"""Incluir conta nova no painel: o MODELO.xlsx, o mapping.yaml e o config.yaml.

Tudo aqui é FICTÍCIO e montado no `tmp_path` — o modelo de verdade tem nome de
empresa e fica fora do repositório. O modelo de teste imita o que importa do
real: fórmulas que citam a própria linha e a aba «Movimentações», os TOTAIS
duas linhas abaixo da última conta, um «Como ler» mesclado embaixo, a
formatação condicional e a área de impressão, e um VLOOKUP de outra aba que
enxerga a faixa das contas.
"""
from datetime import date, datetime
from decimal import Decimal

import openpyxl
import pytest
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill

from conciliacao import painel_novas as pn
from conciliacao.config import load_config
from conciliacao.contas_novas_janela import inclusoes_marcadas
from conciliacao.mapping import AccountMapping
from conciliacao.models import ErpAccount, ErpPayment, Periodo, Snapshot
from conciliacao.pipeline import run_offline

CONFIG = """\
# Configuracao de teste (ficticia). Este comentario tem de sobreviver.
erp:
  api_base: "https://exemplo.invalid"
caminhos:
  modelo: "MODELO.xlsx"
  saida: "output"
  snapshots: "snapshots"
  logs: "logs"
  screenshots: "screenshots"
planilha:
  aba: "Painel"
  primeira_linha: 8
  ultima_linha: 10
  # Duas linhas depois da ultima conta: os totais do dia.
  linha_totais: 12
  ultima_coluna: "M"
  celula_data: "E4"
  celula_total_pagamentos: "D12"
  formato_data: "DD/MM/YYYY"
  colunas_escritas:
    saldo: "C"
    pagamento: "D"
    qtd_sistema: "J"
    qtd_banco: "K"
  colunas_formula: ["E", "F", "G", "I", "L"]
regras:
  excluir_valor_exato: "1.00"
validacao:
  exigir_todos_os_saldos: true
  tolerancia_agregado: "0.01"
"""

MAPA = """\
# Ordem das linhas: comentario que tem de sobreviver.
model_rows:
- row: 8
  label: Conta Alfa - Banco
  uuid: uuid-alfa
  erp_name: ALFA LTDA - BANCO
- row: 9
  label: Conta Beta - Banco
  uuid: uuid-beta
  account_number: 11.111-1
  erp_name: BETA LTDA - BANCO 11.111-1
- row: 10
  label: Conta Gama - Banco
  uuid: uuid-gama
  erp_name: GAMA LTDA - BANCO
# comentario da lista de ignoradas
ignored_erp_accounts:
- Conta Principal
- APENAS LANCAMENTO
"""

ALFA = ErpAccount("uuid-alfa", "ALFA LTDA - BANCO", balance=Decimal("100"))
BETA = ErpAccount("uuid-beta", "BETA LTDA - BANCO 11.111-1", balance=Decimal("50"))
GAMA = ErpAccount("uuid-gama", "GAMA LTDA - BANCO", balance=Decimal("0"))
DELTA = ErpAccount("uuid-delta", "DELTA SPE - BANCO 22.222-2",
                   balance=Decimal("1000"))
SUBCONTA = ErpAccount("uuid-sub",
                      "HOLDING - SUBCONTA 33333-3 - FULANO FICTICIO - BANCO",
                      balance=Decimal("0"))
PESSOA = ErpAccount("uuid-pessoa", "CICLANO  - CARTEIRA", balance=Decimal("10"))
PRINCIPAL = ErpAccount("uuid-principal", "Conta Principal", balance=Decimal("-5"))
CONTAS = [ALFA, BETA, GAMA, DELTA, SUBCONTA, PESSOA, PRINCIPAL]

ROTULOS = ("Conta Alfa - Banco", "Conta Beta - Banco", "Conta Gama - Banco")


def _modelo(caminho):
    wb = openpyxl.Workbook()
    p = wb.active
    p.title = "Painel"
    mov = wb.create_sheet("Movimentações")
    reg = wb.create_sheet("Regras")

    p["B2"] = "CONTROLE"
    p.merge_cells("B2:M2")
    p["B4"] = "Data de Referência:"
    p["B7"] = "CONTA / BANCO"
    for r, rotulo in zip((8, 9, 10), ROTULOS):
        p[f"B{r}"] = rotulo
        p[f"E{r}"] = (f"=SUMIF(Movimentações!$B$9:$B$20,$B{r},"
                      f"Movimentações!$I$9:$I$20)")
        p[f"F{r}"] = (f"=SUMIF(Movimentações!$C$9:$C$20,$B{r},"
                      f"Movimentações!$I$9:$I$20)")
        p[f"G{r}"] = (f'=IF(AND($C{r}="",$D{r}=""),"—",IF($C{r}>=$D{r}+$E{r},'
                      f'"—",$D{r}+$E{r}-$C{r}))')
        p[f"I{r}"] = f'=IF(AND($C{r}="",$D{r}=""),"—",$C{r}-$D{r}-$E{r}+$F{r})'
        # "C10" dentro de aspas é TEXTO: não pode ser tomado por célula.
        p[f"L{r}"] = f'=IF($J{r}="","—","C10 no texto")'
        p[f"M{r}"] = "Aportador Ficticio"
        for c in "BCDEFGHIJKLM":
            p[f"{c}{r}"].font = Font(italic=True)
    p["B12"] = "TOTAIS DO DIA"
    for c in "CDEFGHI":
        p[f"{c}12"] = f"=SUM({c}8:{c}10)"
        p[f"{c}12"].font = Font(bold=True)
    p.row_dimensions[12].height = 21
    p["B14"] = "COMO LER"
    p.merge_cells("B14:D14")
    p["B15"] = "texto de ajuda"
    p.conditional_formatting.add(
        "B8:M10", FormulaRule(formula=['AND($D8=0,$C8<>"")'],
                              fill=PatternFill("solid", start_color="EEEEEE")))
    p.conditional_formatting.add(
        "C8:C10", FormulaRule(formula=["$C8<0"], font=Font(color="FF0000")))
    p.conditional_formatting.add(
        "B2:M2", FormulaRule(formula=["FALSE"], font=Font(color="FF0000")))
    p.print_area = "B2:M12"
    p.freeze_panes = "C8"

    mov["B7"] = "ORIGEM"
    for r in range(9, 12):
        mov[f"G{r}"] = f'=IFERROR(VLOOKUP($C{r},Painel!$B$8:$H$10,7,FALSE),"")'
    mov["G22"] = "=Painel!$D$12"      # o total do painel: desce
    mov["G23"] = "=Painel!$C$9"       # uma conta de cima: fica

    reg["B4"] = "CONTA QUE RECEBE"
    reg["C4"] = "APORTADOR"
    reg["B5"] = "Conta Alfa - Banco"
    reg["C5"] = "Conta Beta - Banco"
    reg["D5"] = 1
    reg["E5"] = "Aporte"
    wb.save(caminho)


@pytest.fixture
def pasta(tmp_path):
    (tmp_path / "config.yaml").write_text(CONFIG, encoding="utf-8")
    (tmp_path / "mapping.yaml").write_text(MAPA, encoding="utf-8")
    _modelo(tmp_path / "MODELO.xlsx")
    return tmp_path


def _bytes(pasta):
    return {nome: (pasta / nome).read_bytes()
            for nome in ("MODELO.xlsx", "mapping.yaml", "config.yaml")}


def _mapa(pasta):
    return AccountMapping.load(pasta / "mapping.yaml")


# --------------------------------------------------------------- achar

def test_fora_do_painel_e_a_mesma_regua_do_aviso(pasta):
    fora = pn.contas_fora_do_painel(CONTAS, _mapa(pasta))
    # A ignorada ("Conta Principal") e as três do painel não aparecem.
    assert [c.id for c in fora] == ["uuid-pessoa", "uuid-delta", "uuid-sub"]


def test_nome_sugerido_desfaz_o_espaco_dobrado():
    assert pn.rotulo_sugerido(PESSOA) == "CICLANO - CARTEIRA"


def test_problemas_vem_todos_de_uma_vez(pasta):
    mapa = _mapa(pasta)
    assert pn.problemas_da_inclusao([], mapa) == ["nenhuma conta marcada."]
    problemas = pn.problemas_da_inclusao([
        pn.Inclusao(ALFA, "Qualquer"),                  # já está no painel
        pn.Inclusao(DELTA, ""),                         # sem nome
        pn.Inclusao(SUBCONTA, "=Subconta"),             # vira fórmula
        pn.Inclusao(PESSOA, "Pessoa * Carteira"),       # curinga no SUMIF
    ], mapa)
    texto = "\n".join(problemas)
    assert "já está no painel (linha 8)" in texto
    assert "falta o nome da linha" in texto
    assert "não pode começar com" in texto
    assert "curinga" in texto


def test_nome_repetido_nao_passa_nem_mudando_a_caixa(pasta):
    """O SUMIF não distingue maiúscula: somaria o aporte nas duas linhas."""
    mapa = _mapa(pasta)
    ja_existe = pn.problemas_da_inclusao([pn.Inclusao(DELTA, "conta alfa - banco")],
                                         mapa)
    assert ja_existe == ["conta alfa - banco: já existe uma linha com esse nome "
                         "(linha 8)."]
    repetido = pn.problemas_da_inclusao([pn.Inclusao(DELTA, "Nova"),
                                         pn.Inclusao(SUBCONTA, "NOVA")], mapa)
    assert repetido == ["NOVA: nome repetido entre as marcadas."]


# --------------------------------------------------------------- o modelo

def test_linhas_novas_copiam_a_ultima_conta_e_o_resto_desce(pasta):
    cfg = load_config(pasta / "config.yaml")
    wb = openpyxl.load_workbook(pasta / "MODELO.xlsx")
    pn.acrescentar_linhas(wb, cfg.planilha, ["Delta SPE", "Holding Subconta"])
    p, mov = wb["Painel"], wb["Movimentações"]

    # As novas: nome na B, fórmulas da linha 10 andando para a linha delas,
    # estilo copiado, e nada do que é dado do dia (C, D, M…).
    assert (p["B11"].value, p["B12"].value) == ("Delta SPE", "Holding Subconta")
    assert p["E11"].value == ("=SUMIF(Movimentações!$B$9:$B$20,$B11,"
                              "Movimentações!$I$9:$I$20)")
    assert p["G12"].value == ('=IF(AND($C12="",$D12=""),"—",IF($C12>=$D12+$E12,'
                              '"—",$D12+$E12-$C12))')
    assert p["L11"].value == '=IF($J11="","—","C10 no texto")'
    assert p["C11"].value is None and p["D11"].value is None
    assert p["M11"].value is None
    assert p["C11"].font.italic

    # A linha em branco antes dos totais continua em branco.
    assert all(p.cell(13, c).value is None for c in range(1, 14))

    # Os totais desceram e somam as novas.
    assert p["B14"].value == "TOTAIS DO DIA"
    assert p["C14"].value == "=SUM(C8:C12)"
    assert p["C14"].font.bold
    assert p.row_dimensions[14].height == 21

    # O «Como ler» desceu com a mesclagem; a do título ficou.
    assert p["B16"].value == "COMO LER"
    mescladas = {m.coord for m in p.merged_cells.ranges}
    assert mescladas == {"B2:M2", "B16:D16"}

    # Formatação condicional e impressão acompanham; a faixa de cima fica.
    faixas = {str(cf.sqref) for cf in p.conditional_formatting}
    assert faixas == {"B8:M12", "C8:C12", "B2:M2"}
    assert str(p.print_area).endswith("$B$2:$M$14")

    # A outra aba enxerga as linhas novas e segue o total que desceu.
    assert mov["G9"].value == '=IFERROR(VLOOKUP($C9,Painel!$B$8:$H$12,7,FALSE),"")'
    assert mov["G22"].value == "=Painel!$D$14"
    assert mov["G23"].value == "=Painel!$C$9"


def test_mapping_ganha_as_linhas_sem_perder_os_comentarios(pasta):
    mapa = _mapa(pasta)
    entradas = pn.entradas_do_mapa(
        [pn.Inclusao(DELTA, 'Delta "SPE": nº #1 Ç'),
         pn.Inclusao(SUBCONTA, "Holding Subconta")], 11, mapa)
    texto = pn.acrescentar_no_mapping(MAPA, entradas)

    assert "# Ordem das linhas: comentario que tem de sobreviver." in texto
    # O bloco novo fica colado no último item, e o comentário continua sendo
    # da chave de baixo.
    assert texto.index("row: 11") < texto.index("# comentario da lista")
    (pasta / "mapping.yaml").write_text(texto, encoding="utf-8")
    novo = _mapa(pasta)
    assert novo.by_row(11).label == 'Delta "SPE": nº #1 Ç'
    assert novo.by_row(11).uuid == "uuid-delta"
    # O número sai do nome quando o ERP não o informa, e vai como texto.
    assert novo.by_row(11).account_number == "222222"
    assert novo.by_row(12).account_number == "333333"
    assert novo.ignored_names == ["Conta Principal", "APENAS LANCAMENTO"]


def test_numero_ja_usado_por_outra_linha_fica_de_fora(pasta):
    """Número repetido no mapa derruba a carga inteira do mapping.yaml."""
    gemea = ErpAccount("uuid-gemea", "OUTRA - BANCO 11.111-1")
    entradas = pn.entradas_do_mapa([pn.Inclusao(gemea, "Outra")], 11, _mapa(pasta))
    assert entradas[0]["account_number"] is None


def test_config_desce_so_as_tres_chaves(pasta):
    texto = pn.atualizar_config(CONFIG, 2)
    assert "Este comentario tem de sobreviver" in texto
    (pasta / "config.yaml").write_text(texto, encoding="utf-8")
    pl = load_config(pasta / "config.yaml").planilha
    assert (pl.ultima_linha, pl.linha_totais, pl.celula_total_pagamentos) == (
        12, 14, "D14")
    assert pl.primeira_linha == 8
    with pytest.raises(pn.InclusaoRecusada):
        pn.atualizar_config(CONFIG.replace("  linha_totais: 12\n", ""), 1)


# --------------------------------------------------------------- incluir

def test_incluir_ponta_a_ponta_e_o_painel_do_dia_sai_com_elas(pasta):
    antes = _bytes(pasta)
    res = pn.incluir_no_painel(
        pasta, [pn.Inclusao(DELTA, "Delta SPE - Banco"),
                pn.Inclusao(SUBCONTA, " Holding Subconta 33333-3 ")],
        CONTAS, agora=datetime(2026, 9, 11, 17, 30, 0))

    assert res.linhas == [(11, "Delta SPE - Banco"),
                          (12, "Holding Subconta 33333-3")]
    # A cópia guarda os três como estavam, e nada de arquivo de trabalho sobra.
    assert res.copia == pasta / "copias do painel" / "2026-09-11 173000"
    assert {nome: (res.copia / nome).read_bytes() for nome in antes} == antes
    sobras = [a.name for a in pasta.iterdir() if ".novo." in a.name
              or ".prova." in a.name]
    assert sobras == []

    # O dia seguinte: o pipeline de sempre, com um pagamento na conta nova.
    cfg = load_config(pasta / "config.yaml")
    mapping = _mapa(pasta)
    hoje = date(2026, 9, 11)
    snapshot = Snapshot(
        reference_date=hoje, collected_at="teste", accounts=CONTAS,
        periodo=Periodo.de_um_dia(hoje),
        payments=[
            ErpPayment(hoje, "Em aberto", Decimal("300"), "Fornecedor Ficticio",
                       "DELTA SPE - BANCO 22.222-2"),
            ErpPayment(hoje, "Em aberto", Decimal("20"), "Outro Ficticio",
                       "ALFA LTDA - BANCO"),
        ])
    resultado = run_offline(snapshot, cfg, mapping)

    assert resultado.classification.unmapped == []
    assert [c.id for c in resultado.balances.contas_desconhecidas] == ["uuid-pessoa"]
    p = openpyxl.load_workbook(resultado.arquivo)["Painel"]
    assert p["B11"].value == "Delta SPE - Banco"
    assert (p["C11"].value, p["D11"].value, p["J11"].value) == (1000, 300, 1)
    assert p["D14"].value == "=SUM(D8:D12)"


def test_recusa_nao_toca_em_nada(pasta):
    antes = _bytes(pasta)
    with pytest.raises(pn.InclusaoRecusada, match="já existe uma linha"):
        pn.incluir_no_painel(pasta, [pn.Inclusao(DELTA, "Conta Beta - Banco")],
                             CONTAS)
    assert _bytes(pasta) == antes
    assert not (pasta / "copias do painel").exists()


def test_nome_parecido_demais_com_outra_conta_e_recusado(pasta):
    """A conta nova não pode puxar para a linha dela uma conta que ninguém
    marcou: os pagamentos da outra entrariam no painel pela linha errada."""
    antes = _bytes(pasta)
    nova = ErpAccount("uuid-zeta", "ZETA LTDA - BANCO")
    vizinha = ErpAccount("uuid-zeta-res", "ZETA LTDA - BANCO - RESERVA")
    with pytest.raises(pn.InclusaoRecusada, match="parecido demais"):
        pn.incluir_no_painel(pasta, [pn.Inclusao(nova, "Zeta")],
                             CONTAS + [nova, vizinha])
    assert _bytes(pasta) == antes


def test_troca_que_falha_no_meio_devolve_o_que_ja_tinha_trocado(pasta, monkeypatch):
    """O Excel com o arquivo aberto: o modelo já foi trocado, o mapa não."""
    antes = _bytes(pasta)
    real = pn.os.replace

    def replace(origem, destino):
        if str(destino).endswith("mapping.yaml"):
            raise PermissionError(13, "arquivo em uso")
        return real(origem, destino)

    monkeypatch.setattr(pn.os, "replace", replace)
    with pytest.raises(pn.InclusaoRecusada, match="Nada mudou"):
        pn.incluir_no_painel(pasta, [pn.Inclusao(DELTA, "Delta")], CONTAS)
    assert _bytes(pasta) == antes
    assert [a.name for a in pasta.iterdir() if ".novo." in a.name] == []


# --------------------------------------------------------------- a janela

def test_a_janela_devolve_so_as_marcadas_com_o_nome_digitado():
    linhas = [[DELTA, True, "  Delta SPE  "], [PESSOA, False, "Pessoa"],
              [SUBCONTA, True, "Holding"]]
    assert inclusoes_marcadas(linhas) == [pn.Inclusao(DELTA, "Delta SPE"),
                                          pn.Inclusao(SUBCONTA, "Holding")]
